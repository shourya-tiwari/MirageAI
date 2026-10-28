"""Build reports/MirageAI_Presentation_V2.pptx, the 16-slide deck for a 30-45 minute
presentation and viva (python-pptx).

    python scripts/build_slides_v2.py

Same design language as scripts/build_slides.py (Calibri, one blue accent,
tinted cards, kicker + title, footer rule), but a longer talk: every slide
carries roughly 2-3 minutes of speaker notes that build each idea from the
basics. The original deck, reports/MirageAI_Slides.pptx, is not touched.

Numbers on slides and in the notes are formatted here from the output tables,
as in build_slides.py. The VIF values come from the T2.6 analysis recorded in
D-023 / D-024 (there is no VIF table).
"""

from __future__ import annotations

import io
import random
import re
import sys
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from mirageai import config as c
from mirageai import models as mo

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "MirageAI_Presentation_V2.pptx"
FIG = ROOT / "outputs" / "figures"
IMG = ROOT / "docs" / "images"

ACCENT = RGBColor(0x2A, 0x78, 0xD6)  # the one accent colour
NAVY = RGBColor(0x14, 0x32, 0x5C)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
RULE = RGBColor(0xE1, 0xE0, 0xD9)
TINT = RGBColor(0xEE, 0xF4, 0xFC)
TINT2 = RGBColor(0xC9, 0xDD, 0xF7)  # highlighted row / held-out blocks
PALE = RGBColor(0xC9, 0xD6, 0xEA)  # light text on navy
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"
W, H = 13.333, 7.5
TOTAL = 16
FOOTER = "MirageAI · Concrete vs. Canopy"
TEAM = ["Shourya Tiwari", "Shashwat Kumar", "Rishik Sinha", "Shubhankar Bhan"]

# ---------------------------------------------------------------------------
# Numbers, all from the tables
# ---------------------------------------------------------------------------
T = lambda n: pd.read_csv(c.table(n))  # noqa: E731
p1 = T("phase8_core_ring_land")
water = T("phase4_water_exclusion").set_index("year")
rr = T("phase4_residual_range")
mc = T("model_comparison").set_index(["scheme", "model"])
gap = T("phase5_cv_gap").set_index(["model", "scheme"])
th = T("phase5_temporal_holdout").set_index(["scheme", "model"])
gi = T("phase6_group_importance").set_index("group")
fi = T("phase6_shap_importance").set_index("feature")
ra = T("phase6_rank_agreement").set_index("comparison")
dg = T("phase6_dominant_driver_groups").set_index("dominant")
ds = T("phase6_dominant_stability").set_index("level")
thr = T("phase6_thresholds").set_index("feature")
rq5 = T("phase6_rq5_frac_built").set_index("subset")
mi = T("phase5_morans_i")
cib = T("phase8_cool_island_bins")
cor = T("phase2_lst_anom_correlations").set_index("Unnamed: 0")
sm = T("phase7_summary")
ens = sm[sm.model == "ensemble"].set_index(["scenario", "bound"])
ms = T("phase7_marginal_summary").set_index("bound")
tc = T("phase7_target_cells")
K = "spatio_temporal_4km"

N = {}
N["ring_lo"] = f"{-p1.core_minus_ring.max():.1f}"
N["ring_hi"] = f"{-p1.core_minus_ring.min():.1f}"
N["rows_all"] = f"{int(water.loc['all', 'rows_before']):,}"
N["cell_years"] = f"{int(water.loc['all', 'rows_after']):,}"
N["cells"] = f"{int(water.loc['2023', 'rows_after']):,}"
N["water_rows"] = f"{int(water.loc['all', 'rows_removed']):,}"
N["water_cells"] = f"{int(water.loc['2023', 'rows_removed']):,}"
N["res_range"] = f"{rr.range_95_m.median():,.0f}"
lo_bin = cib[cib.mid < 10].groupby("year").anom.mean()
hi_bin = cib[(cib.mid > 80) & (cib.mid < 90)].groupby("year").anom.mean()
N["built_gradient"] = f"{(lo_bin - hi_bin).mean():.1f}"
for f, k in [("NDBI", "r_ndbi"), ("frac_grass", "r_grass"), ("frac_tree", "r_tree"), ("NDVI", "r_ndvi"),
             ("frac_built", "r_built")]:
    N[k] = f"{cor.loc[f, 'all_cells']:+.2f}".replace("-", "−")
N["gap_linear"] = f"{gap.loc[('linear', 'spatial_4km'), 'dRMSE']:.2f}"
tree_gaps = [gap.loc[(m, "spatial_4km"), "dRMSE"] for m in ("rf", "xgb", "lgbm")]
N["gap_tree_lo"], N["gap_tree_hi"] = f"{min(tree_gaps):.2f}", f"{max(tree_gaps):.2f}"
N["gap_xgb"] = f"{gap.loc[('xgb', 'spatial_4km'), 'dRMSE']:.2f}"
N["gap8_lgbm"] = f"{gap.loc[('lgbm', 'spatial_8km'), 'dRMSE']:.2f}"
rm = lambda s, m: f"{mc.loc[(s, m), 'RMSE_mean']:.2f}"  # noqa: E731
r2 = lambda s, m: f"{mc.loc[(s, m), 'R2_mean']:.2f}".replace("-", "−")  # noqa: E731
tree_4km = [mc.loc[("spatial_4km", m), "RMSE_mean"] for m in ("rf", "xgb", "lgbm")]
N["tied_lo"], N["tied_hi"] = f"{min(tree_4km):.2f}", f"{max(tree_4km):.2f}"
N["r2_2025"] = f"{th.loc[(K, 'lgbm'), 'R2_unseen_mean']:.2f}"
N["rmse_2025"] = f"{th.loc[(K, 'lgbm'), 'RMSE_unseen_mean']:.2f}"
N["rmse_2025_lin"] = f"{th.loc[(K, 'linear'), 'RMSE_unseen_mean']:.2f}"
N["base_2025"] = f"{th.loc[(K, 'baseline'), 'RMSE_unseen_mean']:.2f}"
mem = {m: th.loc[(K, m), "memorisation_dRMSE_mean"] for m in ("rf", "xgb", "lgbm")}
N["mem_lo"], N["mem_hi"] = f"{min(mem.values()):.2f}", f"{max(mem.values()):.2f}"
N["mem_lgbm"] = f"{mem['lgbm']:.2f}"
N["mem_linear"] = f"{th.loc[(K, 'linear'), 'memorisation_dRMSE_mean']:.2f}"
assert min(mem, key=mem.get) == "lgbm", "slide 10 says LightGBM has the smallest tree memorisation gap"
N["spectral"] = f"{gi.loc['spectral', 'lgbm_all']:.2f}"
N["dry"] = f"{gi.loc['dry cover', 'lgbm_all']:.2f}"
N["tree_grp"] = f"{gi.loc['tree cover', 'lgbm_all']:.2f}"
N["built_grp"] = f"{gi.loc['built', 'lgbm_all']:.2f}"
N["spec_ratio"] = f"{gi.loc['spectral', 'lgbm_all'] / gi.loc['dry cover', 'lgbm_all']:.1f}"
N["ndvi_shap"] = f"{fi.loc['NDVI', 'lgbm_all']:.2f}"
N["ndbi_shap"] = f"{fi.loc['NDBI', 'lgbm_all']:.2f}"
N["mndwi_shap"] = f"{fi.loc['MNDWI', 'lgbm_all']:.2f}"
N["spearman"] = f"{ra.loc['lgbm vs xgb, all cells', 'spearman']:.2f}"
N["dom_spec"] = f"{100 * dg.loc['spectral', 'share']:.1f}"
N["dom_dry"] = f"{100 * dg.loc['dry cover', 'share']:.1f}"
N["group_stab"] = f"{100 * ds.loc['group', 'lgbm_vs_xgb']:.0f}"
t = thr.loc["frac_tree"]
N["knot"] = f"{100 * t.pos_knot:.1f}"
N["rq5_raw"] = f"{rq5.loc['all', 'raw_slope_per_0.1']:.2f}".replace("-", "−")
N["rq5_shap"] = f"{rq5.loc['all', 'shap_slope_per_0.1']:.3f}".replace("-", "−")
N["rq5_shrink"] = f"{100 * (1 - rq5.loc['all', 'shap_slope_per_0.1'] / rq5.loc['all', 'raw_slope_per_0.1']):.0f}"
N["targets"] = f"{len(tc):,}"
lg = mi[(mi.scheme == "spatial_4km") & (mi.model == "lgbm")].morans_I
N["moran"] = f"{lg.mean():.1f}"
for s in ("A", "B", "C"):
    N[f"{s}_lo"] = f"{-ens.loc[(s, 'lower'), 'dLST_excl_first_point']:.2f}"
    N[f"{s}_hi"] = f"{-ens.loc[(s, 'realistic'), 'dLST_excl_first_point']:.2f}"
N["A_lo1"] = f"{-ens.loc[('A', 'lower'), 'dLST_excl_first_point']:.1f}"
N["A_hi1"] = f"{-ens.loc[('A', 'realistic'), 'dLST_excl_first_point']:.1f}"
N["per_pt_below"] = f"{-ms.loc['realistic', 'per_point_2pct_to_knot']:.2f}"
N["per_pt_above"] = f"{-ms.loc['realistic', 'per_point_12_to_30pct']:.2f}"
N["first_pt"] = f"{-ms.loc['realistic', 'first_point_0_to_1pct']:.1f}"
N["n_seeds"] = str(len(c.CV_SEEDS))
N["n_folds"] = str(c.N_SPLITS)
N["n_runs"] = str(len(c.CV_SEEDS) * c.N_SPLITS)
N["n_features"] = str(len(c.FEATURES_TREE))
N["n_iter"] = str(mo.TUNE_N_ITER)
N["subsample"] = f"{mo.TUNE_SUBSAMPLE:,}"
N["block_km"] = str(c.PRIMARY_BLOCK_SIZE_M // 1000)
# T2.6 VIF analysis (D-023, D-024): linear design matrix with frac_built dropped.
VIF = {"NDVI": "29.8", "MNDWI": "18.8", "NDBI": "18.4", "median": "2.2"}

# ---------------------------------------------------------------------------
# Drawing helpers (as in build_slides.py)
# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
cp = prs.core_properties
cp.title = "Concrete vs. Canopy: Explainable Machine Learning for Urban Heat Island Drivers in Pune"
cp.subject = "MirageAI · TE7483 case study, Symbiosis Institute of Technology, Pune · version 2"
cp.author = cp.last_modified_by = "; ".join(TEAM)
cp.keywords = "urban heat island; land surface temperature; spatial cross-validation; SHAP"
cp.comments = ""
BLANK = prs.slide_layouts[6]


def slide_no(slide) -> int:
    return list(prs.slides).index(slide) + 1


def textbox(slide, x, y, w, h, runs, size=18, color=INK2, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, bullet=False, spacing=6, font=FONT):
    """runs: str or list of paragraphs; a paragraph is str or list of (text, overrides) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        parts = [(para, {})] if isinstance(para, str) else para
        if bullet:
            parts = [("▪  ", {"color": ACCENT})] + list(parts)
        for text, ov in parts:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = ov.get("font", font)
            f.size = Pt(ov.get("size", size))
            f.bold = ov.get("bold", bold)
            f.color.rgb = ov.get("color", color)
    return tb


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    box = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.fill.background()
    box.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        box.adjustments[0] = 0.08
    return box


def badge(slide, x, y, d, label, fill=ACCENT, color=WHITE, size=16):
    """A filled circle with a short centred label: the deck's icon motif."""
    rect(slide, x, y, d, d, fill, MSO_SHAPE.OVAL)
    textbox(slide, x - 0.1, y, d + 0.2, d, label, size=size, color=color, bold=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=0)


def arrow(slide, x, y, w=0.4, h=0.36, fill=ACCENT):
    rect(slide, x, y, w, h, fill, MSO_SHAPE.RIGHT_ARROW)


def picture(slide, src, x, y, w, h):
    """Place an image (path or PIL image), scaled to fit the box and centred in it."""
    if isinstance(src, Image.Image):
        buf = io.BytesIO()
        src.save(buf, format="PNG")
        buf.seek(0)
        iw, ih, stream = *src.size, buf
    else:
        path = Path(src) if Path(src).is_absolute() else FIG / src
        iw, ih = Image.open(path).size
        stream = str(path)
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    slide.shapes.add_picture(stream, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                             Inches(pw), Inches(ph))


def new_slide(kicker, title, notes, dark=False):
    s = prs.slides.add_slide(BLANK)
    n = slide_no(s)
    if dark:
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = NAVY
    if title:
        textbox(s, 0.6, 0.62, 12.1, 0.75, title, size=30, color=INK, bold=True)
        textbox(s, 0.62, 0.3, 8, 0.35, kicker.upper(), size=12, color=ACCENT, bold=True)
    rule = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.6), Inches(6.98), Inches(W - 0.6),
                                  Inches(6.98))
    rule.line.color.rgb = RULE
    rule.line.width = Pt(0.75)
    textbox(s, 0.55, 7.03, 6, 0.3, FOOTER, size=10, color=MUTED)
    textbox(s, 11.6, 7.03, 1.2, 0.3, f"{n} / {TOTAL}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = re.sub(r"\s+", " ", notes).strip()
    return s


def callout(slide, x, y, w, big, label, big_size=34, color=ACCENT, label_color=INK2, label_size=15):
    big_h = big_size / 72 * 1.25
    textbox(slide, x, y, w, big_h, big, size=big_size, color=color, bold=True)
    textbox(slide, x, y + big_h, w, 0.9, label, size=label_size, color=label_color)


def card(slide, x, y, w, h, head, body, fill=TINT, head_color=INK, body_color=INK2, size=16):
    rect(slide, x, y, w, h, fill)
    textbox(slide, x + 0.2, y + 0.15, w - 0.4, 0.45, head, size=size + 2, color=head_color, bold=True)
    textbox(slide, x + 0.2, y + 0.62, w - 0.4, h - 0.7, body, size=size, color=body_color)


def table(slide, x, y, widths, rows, header, row_h=0.6, size=16, highlight=None):
    tbl = slide.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(sum(widths)),
                                 Inches(row_h * (len(rows) + 1))).table
    for j, wdt in enumerate(widths):
        tbl.columns[j].width = Inches(wdt)
    for i in range(len(rows) + 1):
        tbl.rows[i].height = Inches(row_h)
        for j in range(len(header)):
            cell = tbl.cell(i, j)
            cell.text = header[j] if i == 0 else rows[i - 1][j]
            para = cell.text_frame.paragraphs[0]
            numeric = j > 0 and re.match(r"^[−\d.]", rows[0][j])
            para.alignment = PP_ALIGN.CENTER if numeric else PP_ALIGN.LEFT
            run = para.runs[0]
            run.font.name, run.font.size = FONT, Pt(size if i else size - 1)
            run.font.bold = i == 0 or j == 0
            run.font.color.rgb = WHITE if i == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = (ACCENT if i == 0 else TINT2 if i == highlight
                                        else TINT if i % 2 else WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.1)
    return tbl


# ---------------------------------------------------------------------------
# 1 Title
# ---------------------------------------------------------------------------
s = new_slide("", None, f"""
Good morning, everyone. We are Shourya Tiwari, Shashwat Kumar, Rishik Sinha and Shubhankar Bhan,
and this is our case study for Applications and Use Cases of Machine Learning at Symbiosis Institute
of Technology, Pune. The project is called MirageAI, and the subtitle is Concrete versus Canopy:
explainable machine learning for urban heat island drivers in Pune.

Let me start with the picture on the right. This is a satellite map of daytime land surface
temperature over Pune and Pimpri-Chinchwad, for the pre-monsoon months of March to May 2024, on a
grid of one-hundred-metre cells. Bright yellow is hot, dark purple is cool. The dark lines are the
rivers and the dark patches in the corner are reservoirs.

Most people expect a city to look like a hot island in a cool countryside. That is the textbook urban
heat island, and it is why we named the project MirageAI. When we looked at the data, the expected
hot core turned out to be a mirage: in the daytime, Pune's dense core is cooler than its dry
surroundings. So the question is no longer "is the city hot?", but "which surfaces are hot, why,
and what would actually cool them?"

Our overall goal has three parts. First, predict how much hotter or cooler each hundred-metre cell is
than the city as a whole, using only open satellite data. Second, explain those predictions, so a
planner can see which factors push a place up or down. Third, use the trained model to run simple
greening scenarios: if we add trees to the hottest urban cells, how much cooler are similar places?

Two principles run through the whole talk. One: we evaluate the model honestly, on places and a year
it has never seen, because ordinary random testing flatters spatial models. Two: every result is an
association learnt from comparing places, not a proof of cause and effect. We will come back to both.

The talk has sixteen slides. We start with the problem and our research questions, then the data and
how we prepared it, then the models, the validation design, the results, explainability with SHAP,
the greening scenarios, and finally applications, limitations and future work.
""")
textbox(s, 0.6, 0.9, 7.0, 1.1, "MirageAI", size=54, color=ACCENT, bold=True)
textbox(s, 0.6, 2.05, 6.6, 1.7, "Concrete vs. Canopy: Explainable Machine Learning for Urban Heat "
        "Island Drivers in Pune", size=28, color=INK, bold=True)
rule = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.62), Inches(4.0), Inches(2.0), Inches(4.0))
rule.line.color.rgb = ACCENT
rule.line.width = Pt(3)
textbox(s, 0.6, 4.2, 6.6, 1.5, [" · ".join(TEAM[:2]), " · ".join(TEAM[2:])], size=17, color=INK, spacing=4)
textbox(s, 0.6, 5.55, 6.8, 0.9, ["Applications and Use Cases of Machine Learning (TE7483)",
                                 "Symbiosis Institute of Technology, Pune"], size=14, color=MUTED, spacing=2)
picture(s, "01_lst_map_2024.png", 7.45, 0.55, 5.5, 6.25)

# ---------------------------------------------------------------------------
# 2 Problem statement
# ---------------------------------------------------------------------------
s = new_slide("Problem statement", "Cities are heating, and planners need to know where to act", f"""
Let me begin with the basic idea. An urban heat island is a part of a city that is hotter than its
surroundings. It happens because we replace soil, grass and trees with roofs, roads and paving. These
surfaces absorb sunlight and store heat, and they have little or no water to evaporate. Vegetation,
by contrast, cools itself the way we sweat: through evapotranspiration, it uses solar energy to turn
water into vapour instead of warming the surface. Trees also shade the ground.

There are two kinds of heat island. The air-temperature heat island is what people feel. The surface
heat island is the skin temperature of the ground and roofs, which satellites measure with
thermal-infrared sensors. Our project works with the surface version, called land surface
temperature, or LST.

Why Pune? Pune and Pimpri-Chinchwad are among the fastest-growing urban areas in India. Built-up land
keeps spreading outward into what used to be farmland, grassland and scrub. In the pre-monsoon months,
the city-wide median surface temperature in our data is above forty-four degrees Celsius. Heat
stresses health, raises electricity demand for cooling, and hits outdoor workers hardest. Indian
cities are now writing heat action plans, and the practical question in those plans is: where would
greening cool the most?

Here is the twist on the right of the slide. In daytime Pune, land cells within five kilometres of
the centre are {N['ring_lo']} to {N['ring_hi']} °C cooler than the outer ring beyond twelve
kilometres, in every one of our three years. So if a planner assumes a hot core, they may plant trees
in the wrong places.

Why is previous work not enough? Many studies correlate temperature with one index, like NDVI, or fit a
linear model, which cannot capture thresholds or interactions. And most machine-learning studies test
their models on randomly chosen pixels. Neighbouring pixels are nearly identical, so the model is
tested on near-copies of its training data, and the accuracy looks better than it really is.

That is why explainable AI matters here. A planner will not act on a black box. They need to know why
a place is hot, which factors matter, and by roughly how much, and they need to trust that the
explanation is stable. So our problem is: predict surface heat accurately in places the model has never
seen, explain it, and turn it into actionable, honest guidance.
""")
probs = [("Rapid urbanisation", "built-up land spreading outward"),
         ("Urban Heat Island", "hard surfaces store and re-emit heat"),
         ("Loss of vegetation", "canopy replaced by roofs and paving"),
         ("Rising surface heat", "pre-monsoon city median above 44 °C"),
         ("Explainable prediction", "planners will not act on a black box"),
         ("Actionable planning", "where to green, and by how much")]
for i, (a, b) in enumerate(probs):
    row, col = divmod(i, 2)
    x0, y0 = 0.6 + col * 3.95, 1.6 + row * 1.75
    rect(s, x0, y0, 3.75, 1.5, TINT)
    badge(s, x0 + 0.2, y0 + 0.22, 0.55, str(i + 1))
    textbox(s, x0 + 0.9, y0 + 0.2, 2.75, 0.55, a, size=17, color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE,
            spacing=0)
    textbox(s, x0 + 0.2, y0 + 0.9, 3.4, 0.55, b, size=15, spacing=0)
rect(s, 8.65, 1.6, 4.05, 5.0, NAVY)
textbox(s, 8.95, 1.85, 3.6, 0.4, "PUNE'S TWIST", size=13, color=PALE, bold=True)
callout(s, 8.95, 2.3, 3.6, f"{N['ring_lo']}–{N['ring_hi']} °C", "cooler: city core vs outer ring, "
        "daytime, 2023–25", big_size=40, color=WHITE, label_color=PALE)
textbox(s, 8.95, 3.9, 3.6, 2.5, ["Assume a hot core and you green the wrong places.",
                                  "Random-pixel testing flatters models."], size=16, color=WHITE,
        bullet=True, spacing=10)

# ---------------------------------------------------------------------------
# 3 Research objectives and research questions
# ---------------------------------------------------------------------------
s = new_slide("Research objectives", "Research objectives and research questions", f"""
This slide sets out what we promised to deliver. On the left are five objectives; on the right are
five research questions, each of which gets a measured answer by the end of the talk.

Objective one is prediction: estimate each cell's land surface temperature anomaly, meaning how much
hotter or cooler it is than that year's city-wide median, from satellite-derived features alone.
Objective two is to identify the drivers: which kinds of surface go with heat, and which with cooling.
Objective three is scientific reliability. Spatial data break the usual machine-learning assumption that
samples are independent, so we need a validation design that respects geography. Objective four is
explanation: every prediction should be decomposable into feature contributions, using SHAP. Objective
five is decision support: use the model to simulate greening and give planners a range, not a single
over-confident number.

Now the research questions. RQ1 asks whether satellite variables can accurately predict LST anomalies
across Pune. Accuracy here means accuracy in unseen places and an unseen year, not on neighbours of
the training data. RQ2 asks which environmental factors contribute most: vegetation, dryness, moisture,
land cover or terrain. This is the "concrete versus canopy" question in our title.

RQ3 compares spatial block cross-validation with a conventional random train-test split. This matters
because if random testing overstates accuracy, the whole literature that relies on it is affected, and
it could even change which model we pick.

RQ4 asks whether SHAP can interpret the model reliably. An explanation is only useful if it is stable,
so we check it across two different models and two SHAP methods, and we look at where it is not stable.

RQ5 asks whether counterfactual, or what-if, simulations can estimate the cooling associated with more
urban vegetation. Here we must be careful: the model learns associations between places, so a
simulation says how much cooler similar places with more trees are, not a guaranteed effect of planting.

Each research question maps to a part of the pipeline: RQ1 and RQ3 to evaluation, RQ2 and RQ4 to
explainability, and RQ5 to the scenario analysis. With that, let us look at the data.
""")
textbox(s, 0.6, 1.55, 5.0, 0.45, "Objectives", size=20, color=ACCENT, bold=True)
textbox(s, 0.6, 2.1, 5.0, 4.6, ["Predict LST anomaly from satellite features",
                                "Identify the strongest drivers of urban heat",
                                "Build a spatially reliable ML framework",
                                "Explain predictions with SHAP",
                                "Simulate greening to support planning"], size=18, bullet=True, spacing=26)
rqs = ["Can satellite features accurately predict LST anomalies across Pune?",
       "Which environmental factors contribute most to urban heat?",
       "How does spatial block CV compare with a random train–test split?",
       "Can SHAP reliably interpret the model's predictions?",
       "Can counterfactual simulations estimate cooling from more vegetation?"]
for i, q in enumerate(rqs):
    y0 = 1.6 + i * 1.03
    rect(s, 5.95, y0, 1.0, 0.86, ACCENT)
    textbox(s, 5.95, y0, 1.0, 0.86, f"RQ{i + 1}", size=19, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, spacing=0)
    rect(s, 7.05, y0, 5.65, 0.86, TINT)
    textbox(s, 7.2, y0, 5.4, 0.86, q, size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE, spacing=0)

# ---------------------------------------------------------------------------
# 4 Datasets
# ---------------------------------------------------------------------------
s = new_slide("Datasets", "Four open satellite datasets, one 100 m grid", f"""
We use four open datasets, all accessed through Google Earth Engine, which lets us process years of
satellite imagery in the cloud without downloading raw scenes. The study area is a box over Pune and
Pimpri-Chinchwad, and the period is the pre-monsoon season, March to May, of 2023, 2024 and 2025. We
chose pre-monsoon because skies are clear and surface heat peaks, so the signal is strong and few
images are lost to cloud.

The first dataset is Landsat 8 and 9, Collection 2 Level 2. Their thermal sensor measures emitted
infrared radiation, from which the surface temperature product is derived. This is our target, the
thing we predict. The thermal sensor's native resolution is about one hundred metres, which is why our
whole grid is one hundred metres.

The second is Sentinel-2 surface reflectance, at ten to twenty metres. Reflectance in the red, near-
infrared, green and short-wave infrared bands gives us three spectral indices: NDVI for greenness,
NDBI, which turned out to measure dryness, and MNDWI for moisture. The map on the left is NDVI for 2024:
green is vegetated, and you can see the river corridors and the green hills.

The third is ESA WorldCover 2021, a global land-cover map at ten metres. Each ten-metre pixel has a
class such as tree, grass, crop, built-up or water. We turn these into fractions: what share of each
hundred-metre cell is trees, what share is built-up, and so on. We also compute distance to water from it.

The fourth is the SRTM digital elevation model at thirty metres, which gives elevation and slope.
Pune has hills, and higher, steeper, drier slopes behave differently from the river plain.

Why combine four sources? Because no single sensor sees all the drivers. Landsat sees heat but not why.
Sentinel-2 sees the current state of vegetation and moisture. WorldCover sees what the surface is made
of, and SRTM sees the shape of the land. Together they give {N['n_features']} features per cell, for
{N['cells']} land cells per year, which is {N['cell_years']} cell-years in total. One important rule:
the target comes only from Landsat, and no feature is derived from it, so there is no target leakage.
""")
picture(s, "01_ndvi_map_2024.png", 0.5, 1.5, 6.3, 5.3)
tiles = [("Landsat 8/9", "100 m thermal", "target: land surface temperature"),
         ("Sentinel-2", "10–20 m", "NDVI, NDBI, MNDWI indices"),
         ("ESA WorldCover 2021", "10 m", "7 land-cover fractions, distance to water"),
         ("SRTM DEM", "30 m", "elevation, slope")]
for i, (a, res, b) in enumerate(tiles):
    y0 = 1.55 + i * 1.1
    rect(s, 7.2, y0, 5.5, 0.95, TINT)
    textbox(s, 7.4, y0 + 0.08, 5.2, 0.85, [[(a, {"bold": True, "color": INK}), (f"   {res}", {"color": MUTED})],
                                           b], size=16, spacing=2)
textbox(s, 7.2, 6.0, 5.6, 0.8, f"Mar–May 2023, 2024, 2025 · {N['cell_years']} land cell-years",
        size=17, color=ACCENT, bold=True)

# ---------------------------------------------------------------------------
# 5 Data preprocessing
# ---------------------------------------------------------------------------
s = new_slide("Data preprocessing", "From raw scenes to one clean modelling table", f"""
Raw satellite images cannot go straight into a model. They come from different sensors, in different
resolutions, with clouds and gaps. This slide walks through the six preprocessing steps in order.

Step one, projection. Satellite products arrive in latitude and longitude, where a degree is not a fixed
distance. We reproject everything to UTM zone 43 North, code EPSG 32643, where units are metres. That
makes areas, slopes, distances and our four-kilometre blocks exact.

Step two, resampling. Our sources are ten, twenty and thirty metres. We first take a seasonal median of
each pixel over March to May, which removes passing clouds and odd days. Then we aggregate to one hundred
metres with an area-weighted mean, so each cell is the average of the finer pixels inside it. For land
cover, we first convert each class to a zero-or-one mask and then average, so the result is the fraction
of the cell in that class.

Step three, the one-hundred-metre grid itself, which matches the Landsat thermal sensor. Modelling at
thirty metres would pretend to detail the target does not have. Each cell gets a stable identifier
computed from its coordinates, so the same cell in 2023, 2024 and 2025 always falls in the same block.

Step four, the water mask. Cells that are at least half water, the reservoirs and wide river reaches,
are removed: {N['water_cells']} cells per year, {N['water_rows']} of {N['rows_all']} rows, none of them
urban. Water temperature is governed by thermal inertia, not land-surface drivers, and water is not
something a planner can change.

Step five, missing values. Landsat and Sentinel-2 come with quality flags; we mask cloud, cirrus and
cloud-shadow pixels and skip scenes over thirty percent cloud. Masked pixels get an explicit no-data value
of minus 9999, so a gap can never be mistaken for a real zero. We require at least {c.MIN_LST_OBS} clear
observations per cell. Thanks to the clear pre-monsoon season, no cell failed that check.

Step six, the target. Absolute temperature shifts between years with weather and the timing of the
overpass, and nothing in our features knows the year. So we model the anomaly: each cell's LST minus that
year's city-wide median. A positive value means hotter than typical, negative means cooler.
""")
steps = [("Projection", "UTM 43N (EPSG:32643), metres"),
         ("Resampling", "seasonal median, area-weighted mean"),
         ("100 m grid", f"Landsat thermal scale · {N['cells']} land cells"),
         ("Water mask", f"≥ 50% water removed · {N['water_rows']} rows"),
         ("Missing values", "cloud masks · nodata −9999 · none dropped"),
         ("LST anomaly", "cell LST − that year's median")]
for i, (a, b) in enumerate(steps):
    row, col = divmod(i, 3)
    x0, y0 = 0.6 + col * 4.15, 1.6 + row * 2.2
    rect(s, x0, y0, 3.65, 1.85, TINT)
    badge(s, x0 + 0.2, y0 + 0.2, 0.55, str(i + 1))
    textbox(s, x0 + 0.9, y0 + 0.25, 2.65, 0.5, a, size=18, color=INK, bold=True, spacing=0)
    textbox(s, x0 + 0.2, y0 + 0.95, 3.3, 0.85, b, size=15, spacing=0)
    if col < 2:
        arrow(s, x0 + 3.71, y0 + 0.75)
textbox(s, 0.6, 6.15, 12.1, 0.6, [[("Why the anomaly?  ", {"color": ACCENT, "bold": True}),
                                  ("the yearly level shifts with weather; the model learns the spatial "
                                   "pattern, not the year.", {})]], size=17)

# ---------------------------------------------------------------------------
# 6 Feature engineering
# ---------------------------------------------------------------------------
s = new_slide("Feature engineering", f"{N['n_features']} physically meaningful features per cell", f"""
Feature engineering means turning raw measurements into inputs that carry meaning for the problem. A
model given raw reflectance in twelve bands has to discover vegetation by itself. A model given NDVI
starts from a quantity that physically means greenness. Good features make models more accurate, easier
to explain, and more robust in new places. We built {N['n_features']} features in three families.

First, spectral indices from Sentinel-2. Each one is a normalised difference of two bands, a ratio
between minus one and plus one. NDVI compares near-infrared and red: healthy leaves reflect strongly in
near-infrared and absorb red, so high NDVI means dense, green vegetation. NDBI compares short-wave
infrared and near-infrared. It was designed to highlight built-up land, but as we will see on the next
slide, in Pune it behaves as a dryness index. MNDWI compares green and short-wave infrared and responds to
water and moisture, so we call it the moisture index.

Second, land-cover fractions from WorldCover: the share of each cell that is trees, grass, crops, water,
bare land, shrubs and built-up. Fractions are better than a single class label, because a cell that is
sixty percent built and forty percent trees is very different from one that is one hundred percent
built. The seven fractions sum to one. That creates perfect collinearity for linear regression, so for
the linear model only, we drop the built-up fraction as the reference category. The tree models keep all
seven, otherwise SHAP could never attribute anything to concrete.

Third, terrain: elevation and slope from SRTM, and distance to the nearest water body in metres, which
captures the cooling influence of rivers and reservoirs.

Just as important is what we left out. Coordinates, latitude, longitude, easting and northing, are never
features. If they were, the model could memorise that a certain location is hot without learning why,
and its explanations would say "it is hot because of where it is", which is useless to a planner. The
code asserts this rule, along with excluding the year and the observation count.
""")
fams = [("S", "Spectral", [("NDVI", "greenness · B8, B4"), ("NDBI", "dryness · B11, B8"),
                           ("MNDWI", "moisture · B3, B11")]),
        ("L", "Land-cover fractions", [("Trees · Grass · Crops", ""), ("Water · Bare land", ""),
                                       ("Shrubs · Built-up", ""), ("share of cell, sum = 1", None)]),
        ("T", "Terrain", [("Elevation", "SRTM, metres"), ("Slope", "degrees"),
                          ("Distance to water", "metres")])]
for i, (ic, head, items) in enumerate(fams):
    x0 = 0.6 + i * 4.15
    rect(s, x0, 1.6, 3.85, 4.3, TINT)
    badge(s, x0 + 0.25, 1.85, 0.6, ic, size=18)
    textbox(s, x0 + 1.0, 1.9, 2.75, 0.55, head, size=19, color=INK, bold=True, spacing=0)
    paras = []
    for a, b in items:
        if b is None:
            paras.append([(a, {"color": MUTED, "size": 15})])
        elif b:
            paras.append([(a, {"bold": True, "color": INK}), (f"  {b}", {})])
        else:
            paras.append([(a, {"bold": True, "color": INK})])
    textbox(s, x0 + 0.25, 2.75, 3.45, 3.0, paras, size=16, spacing=14)
    why = ["vegetation and water state, every year", "what the surface is made of",
           "shape of the land, nearness of water"][i]
    textbox(s, x0 + 0.25, 4.95, 3.45, 0.8, [[("Captures:  ", {"bold": True, "color": ACCENT}), (why, {})]],
            size=15, spacing=0)
textbox(s, 0.6, 6.15, 12.1, 0.6, [[("Never features:  ", {"color": ACCENT, "bold": True}),
                                  ("coordinates, year, observation count.  ", {"color": INK}),
                                  ("Linear model drops built-up (sum-to-1).", {})]], size=17)

# ---------------------------------------------------------------------------
# 7 Exploratory data analysis
# ---------------------------------------------------------------------------
s = new_slide("Exploratory data analysis", "What the data said before any model", f"""
Before training any model, we looked hard at the data. Exploratory analysis tells us what relationships
exist, which assumptions are wrong, and which features might cause trouble.

The left panel is the correlation matrix of all features and the target, pooled over the three years.
Red is positive, blue negative. Look at the last row, the LST anomaly. The strongest positive
correlations are NDBI at {N['r_ndbi']} and grass fraction at {N['r_grass']}: dry grassland is the hottest
surface in Pune. The strongest negative ones are tree fraction at {N['r_tree']} and NDVI at
{N['r_ndvi']}: vegetation goes with cooler surfaces. And surprisingly, built-up fraction is negative,
{N['r_built']}. More concrete goes with slightly cooler daytime surfaces.

The feature histograms, bottom right, show the shape of each variable. The land-cover fractions are
zero-inflated: most cells have none of a given class, with a spike at zero and sometimes at one. The
indices are roughly bell-shaped. These shapes favour tree models, which split on thresholds and do not
assume normality.

The scatter-style panels, top right, bin cells by their built-up fraction. On the left, mean temperature
anomaly falls as built-up share rises: the least-built cells run {N['built_gradient']} °C warmer than
cells that are eighty to ninety percent built, in all three years. On the right is the important
discovery. NDBI, the so-called built-up index, is highest in the least-built cells. Why? Because NDBI
responds to short-wave infrared, which is high for dry grass, fallow fields and bare soil. So in Pune,
NDBI behaves more like a dryness indicator than a pure built-up indicator, and we call it that throughout.

We also checked multicollinearity with the variance inflation factor, or VIF. A VIF above ten means a
feature is largely predictable from the others. With built-up dropped, the median VIF is
{VIF['median']}, but the three indices are high: NDVI {VIF['NDVI']}, MNDWI {VIF['MNDWI']} and NDBI
{VIF['NDBI']}. They overlap, so in the linear model their individual coefficients are unstable, and when
we explain the tree models we rank them together as one spectral group.
""")
picture(s, "02_correlation_heatmap.png", 0.4, 1.45, 4.9, 4.4)
textbox(s, 0.6, 5.9, 4.6, 0.9, [[("Correlation with LST anomaly", {"bold": True, "color": INK})],
                                [(f"NDBI {N['r_ndbi']} · grass {N['r_grass']} · trees {N['r_tree']}", {})]],
        size=15, spacing=2)
picture(s, "08_cool_island.png", 5.45, 1.45, 7.5, 3.0)
picture(s, "02_feature_distributions.png", 5.45, 4.6, 2.95, 2.25)
rect(s, 8.6, 4.6, 4.1, 2.2, TINT)
textbox(s, 8.8, 4.7, 3.8, 0.4, "Discovery", size=17, color=ACCENT, bold=True, spacing=0)
textbox(s, 8.8, 5.1, 3.8, 0.85, "NDBI is highest in the least-built cells: it tracks dryness, "
        "not concrete.", size=15, color=INK, spacing=0)
textbox(s, 8.8, 6.0, 3.8, 0.75, f"VIF: NDVI {VIF['NDVI']} · MNDWI {VIF['MNDWI']} · NDBI {VIF['NDBI']}",
        size=14, spacing=0)

# ---------------------------------------------------------------------------
# 8 Machine learning models
# ---------------------------------------------------------------------------
s = new_slide("Machine learning models", "A baseline and four learners on identical folds", f"""
We compared five models, from the simplest possible to state-of-the-art gradient boosting, all trained
and tested on exactly the same folds, so the comparison is fair.

The mean baseline is not really a model. It predicts the training-set average for every cell. Its job is
to set the floor: any useful model must beat it clearly. Because our target is an anomaly around zero,
the baseline's error is simply the spread of the data.

Linear regression predicts the target as a weighted sum of the features, plus a constant. It is fast and
easy to read: each coefficient says how much the prediction changes per unit of a feature. But it assumes
straight-line, additive effects. It cannot learn that tree cover helps a lot at first and then levels
off, or that dryness matters more on slopes. We standardise features inside each training fold, so test
data never leaks into the scaling. Linear regression represents the status quo in many heat island
studies.

Random Forest builds hundreds of decision trees. Each tree is grown on a bootstrap sample of the rows
and considers random subsets of features at each split, so the trees are different from each other.
The forest averages them. Averaging reduces variance, so the forest is robust and captures
nonlinearity and interactions automatically. This idea is called bagging.

XGBoost and LightGBM use the other big idea, boosting. Instead of independent trees, they add small
trees one at a time, and each new tree is fitted to the errors left by the trees so far, with a small
learning rate so the model improves gradually. XGBoost grows each tree level by level, depth-wise.
LightGBM grows leaf-wise, always splitting the leaf that reduces error most, and it bins feature values
into histograms, which makes it very fast on our quarter of a million rows.

All hyperparameters were tuned by randomised search, {N['n_iter']} candidate settings per model, scored
inside the {N['block_km']}-kilometre spatial folds, on 2023 and 2024 only, using a {N['subsample']}-row
sample. So the 2025 test year never influenced a single setting, and no tuning choice saw its own test
blocks.
""")
table(s, 0.6, 1.6, [2.4, 2.5, 3.95, 3.25],
      [("Mean baseline", "constant", "predicts the training mean", "the floor to beat"),
       ("Linear Regression", "linear", "weighted sum of features", "interpretable status quo"),
       ("Random Forest", "bagging", "averages many deep trees", "robust, nonlinear"),
       ("XGBoost", "boosting, depth-wise", "each tree fixes past errors", "strong on tabular data"),
       ("LightGBM", "boosting, leaf-wise", "histogram splits, leaf-wise growth", "fast, fast exact SHAP")],
      ["Model", "Family", "How it learns", "Why include it"], row_h=0.62, size=16)
rect(s, 0.6, 5.65, 12.1, 1.1, TINT)
textbox(s, 0.85, 5.72, 11.6, 0.45, "Tuning, inside spatial CV", size=18, color=INK, bold=True, spacing=0)
textbox(s, 0.85, 6.15, 11.6, 0.5, f"Randomised search · {N['n_iter']} candidates per model · "
        f"{N['block_km']} km spatial folds · 2023–24 only · {N['subsample']}-row sample", size=16, spacing=0)

# ---------------------------------------------------------------------------
# 9 Spatial cross-validation
# ---------------------------------------------------------------------------
s = new_slide("Spatial cross-validation", "Random split vs spatial block split", f"""
This is the methodological heart of the project, so let me build it up carefully.

Machine-learning evaluation assumes the test data are independent of the training data. Geographic data
break that assumption. Tobler's first law of geography says near things are more related than distant
things. Two neighbouring hundred-metre cells usually have almost the same land cover, vegetation and
temperature. In statistics this is called spatial autocorrelation.

Now look at the left diagram, a random split. We pick cells at random for testing, shown in blue. Almost
every test cell has training cells right next to it. A flexible model can effectively look up the
answer from a neighbour it has already seen. This is spatial leakage, and it makes accuracy look better
than it will be in a genuinely new place.

The right diagram is spatial block cross-validation. We divide the city into square blocks and hold out
whole blocks at once. Now the test cells are surrounded by other test cells, and the model must predict
from what a place is like, not from what its neighbours were. This is the honest question a planner
asks: how good is the model in a neighbourhood it has not seen?

How big should a block be? Big enough that the remaining errors are no longer correlated across its
edge. We fitted a variogram to the model's residuals, and the correlation range came out at about
{N['res_range']} metres, so we use {N['block_km']}-kilometre blocks. Blocks are dealt to {N['n_folds']}
folds at random, repeated over {N['n_seeds']} seeds, so every score is a mean of {N['n_runs']} runs. We
add a spatio-temporal test, too: train on 2023 and 2024, test on 2025 in unseen blocks.

What did it show? Moving from random to {N['block_km']}-kilometre blocks, linear regression's error rises by
only {N['gap_linear']} °C, but the tree models rise by {N['gap_tree_lo']} to {N['gap_tree_hi']} °C. The more
flexible a model, the more random splitting flatters it, so random testing does not just inflate scores,
it can change which model looks best. To be precise, spatial CV does not make the model better. It
measures generalisation to new places honestly, which is what lets us trust everything that follows.
""")
rnd = random.Random(c.SEED)
G, cs = 12, 0.27  # 12 x 12 cells, 0.27" each
panels = [(0.6, "Random split", "test cells touch training neighbours → leakage, optimistic score"),
          (4.95, "Spatial block split", f"whole {N['block_km']} km blocks held out → honest score")]
test_blocks = {(0, 1), (1, 3), (2, 0), (3, 2)}
for k, (x0, head, cap) in enumerate(panels):
    textbox(s, x0, 1.5, 3.9, 0.45, head, size=19, color=INK, bold=True, spacing=0)
    for r in range(G):
        for q in range(G):
            held = rnd.random() < 0.25 if k == 0 else (r // 3, q // 3) in test_blocks
            gx = 0.05 * (q // 3) if k == 1 else 0
            gy = 0.05 * (r // 3) if k == 1 else 0
            rect(s, x0 + q * cs + gx, 2.05 + r * cs + gy, cs - 0.03, cs - 0.03,
                 ACCENT if held else TINT2, MSO_SHAPE.RECTANGLE)
    textbox(s, x0, 5.45, 3.6, 0.85, cap, size=15, color=INK2, spacing=0)
arrow(s, 4.3, 3.55, 0.5, 0.45)
rect(s, 0.6, 6.4, 0.22, 0.22, ACCENT, MSO_SHAPE.RECTANGLE)
textbox(s, 0.88, 6.33, 1.2, 0.35, "test", size=14, spacing=0)
rect(s, 1.75, 6.4, 0.22, 0.22, TINT2, MSO_SHAPE.RECTANGLE)
textbox(s, 2.03, 6.33, 1.2, 0.35, "train", size=14, spacing=0)
stats = [(f"{N['res_range']} m → {N['block_km']} km", "residual autocorrelation range sets the block"),
         (f"{N['n_seeds']} × {N['n_folds']} = {N['n_runs']} runs", "seeds × folds behind every score"),
         (f"+{N['gap_linear']} vs +{N['gap_tree_lo']}–{N['gap_tree_hi']} °C",
          "RMSE inflation, random → blocks: linear vs trees")]
for i, (big, lab) in enumerate(stats):
    y0 = 1.6 + i * 1.75
    rect(s, 9.15, y0, 3.55, 1.55, TINT)
    callout(s, 9.35, y0 + 0.1, 3.25, big, lab, big_size=24, label_size=14)

# ---------------------------------------------------------------------------
# 10 Results
# ---------------------------------------------------------------------------
s = new_slide("Results", "Honest accuracy: the tree models tie, LightGBM is selected", f"""
Here are the numbers. The main metric is RMSE, root mean squared error, in degrees Celsius: roughly, the
typical size of a prediction error, with large errors penalised more. Lower is better. R squared is the
share of the variation in the target that the model explains; one is perfect and zero is no better than
the mean.

Read the table column by column. Under random cross-validation, XGBoost looks best at
{rm('random', 'xgb')} °C. That is the number many studies would publish. Under {N['block_km']}-kilometre
spatial blocks, everything gets worse, and the three tree models are tied between {N['tied_lo']} and
{N['tied_hi']} °C, well inside their fold-to-fold spread. So the random-CV winner was partly an artefact
of leakage.

The last two columns are the toughest test: 2025, a year no model saw during training or tuning, in blocks
it never saw. LightGBM scores {N['rmse_2025']} °C with R squared {N['r2_2025']}. Random Forest and XGBoost
are within a few thousandths of a degree, so honestly it is a tie, not a decisive win.

Compare with the floor. The mean baseline's error is about {N['base_2025']} °C, so the tree models roughly
halve it. Linear regression reaches {N['rmse_2025_lin']} °C: decent, which tells us much of the signal is
fairly simple, but the trees add real value by capturing thresholds such as the tree-cover knot we will
see later.

So why choose LightGBM as the final model? Three reasons. One, it is tied for the best accuracy on the
honest tests and marginally lowest on 2025. Two, it has the smallest memorisation gap of the three tree
models: when it has already seen a place in earlier years, it improves by {N['mem_lgbm']} °C, against up
to {N['mem_hi']} for Random Forest. A smaller gap means it relies less on memorising places. Three, it is
fast to train and has exact, fast TreeSHAP explanations, which matters when we explain tens of thousands of
cells. We still use the others as a check: every explanation is repeated with XGBoost, and every scenario
is averaged across all three tree models.
""")
mrows = []
for m, name in [("baseline", "Mean baseline"), ("linear", "Linear Regression"), ("rf", "Random Forest"),
                ("xgb", "XGBoost"), ("lgbm", "LightGBM")]:
    mrows.append((name, rm("random", m), rm("spatial_4km", m), r2("spatial_4km", m),
                  f"{th.loc[(K, m), 'RMSE_unseen_mean']:.2f}",
                  f"{th.loc[(K, m), 'R2_unseen_mean']:.2f}".replace("-", "−")))
table(s, 0.6, 1.6, [2.15, 1.45, 1.3, 1.15, 1.4, 1.15], mrows,
      ["Model", "RMSE random", "RMSE 4 km", "R² 4 km", "RMSE 2025", "R² 2025"], row_h=0.66, size=16,
      highlight=5)
textbox(s, 0.6, 5.7, 8.6, 0.9, ["RMSE in °C, mean of 25 runs; lower is better.",
                                "2025 = unseen year and unseen 4 km blocks."], size=14, color=MUTED, spacing=2)
callout(s, 9.4, 1.55, 3.4, f"{N['rmse_2025']} °C", "RMSE, 2025, unseen places", big_size=30)
callout(s, 9.4, 2.95, 3.4, f"R² {N['r2_2025']}", "variance explained, 2025", big_size=30)
rect(s, 9.4, 4.4, 3.3, 2.35, ACCENT)
textbox(s, 9.6, 4.52, 3.0, 0.45, "LightGBM selected", size=18, color=WHITE, bold=True, spacing=0)
textbox(s, 9.6, 5.0, 3.0, 1.7, ["Tied-best accuracy", "Least memorisation of the trees",
                                "Fast, exact TreeSHAP"], size=15, color=WHITE, spacing=6)

# ---------------------------------------------------------------------------
# 11 Explainable AI
# ---------------------------------------------------------------------------
local = Image.open(IMG / "dashboard_interactive_prediction.png").convert("RGB").crop((595, 795, 1135, 1142))
s = new_slide("Explainable AI", "SHAP: what pushes each prediction up or down", f"""
A good prediction is not enough; we need to know why. For this we use SHAP, which stands for SHapley
Additive exPlanations.

The idea comes from cooperative game theory. Imagine a team of four people wins a prize, and we want to
split it fairly according to each member's contribution. The problem is that contributions depend on who
else is in the team. Lloyd Shapley's solution, which later contributed to a Nobel Prize, is to consider
every possible order in which members could join the team, measure how much each member adds when they
join, and average that across all orders. That average is the member's Shapley value.

In SHAP, the features are the team members and the prediction is the prize. For one cell, SHAP starts from
the model's average prediction, the base value, and assigns each feature a contribution in degrees
Celsius, so that the base value plus all contributions equals the cell's prediction exactly. That
additivity is what makes SHAP so readable.

That gives us local explanations, shown on the right: one hot, tree-less fringe cell. Red bars push it
warmer: high dryness, no tree cover, low NDVI. Blue bars push it cooler. A planner can read exactly why
this cell is predicted hot.

Collect SHAP values over ten thousand cells and you get the global view, on the left: a beeswarm plot. Each
dot is one cell; its position is the feature's contribution, and its colour is the feature's value, dark
for high. High NDVI, dark dots, sit on the left: greener means cooler. High dryness sits on the right.
These patterns are also our contribution plots. Feature importance is the mean absolute SHAP value: NDVI
{N['ndvi_shap']}, NDBI {N['ndbi_shap']} and MNDWI {N['mndwi_shap']} °C lead.

Is it reliable? The three indices are correlated, so the model can trade credit between them. We therefore
rank feature groups. The spectral group leads at {N['spectral']} °C, about {N['spec_ratio']} times the
dry-cover group, and built-up is the smallest group, {N['built_grp']} °C. XGBoost gives the same ranking,
with rank correlation {N['spearman']}, and the same leading group in {N['group_stab']}% of cells. Remember
that SHAP explains the model, not the physical world: these are associations.
""")
textbox(s, 0.6, 1.5, 6.0, 0.4, "GLOBAL SHAP · feature contributions, 10,000 cells", size=13, color=ACCENT,
        bold=True, spacing=0)
picture(s, "06_shap_beeswarm.png", 0.45, 1.9, 6.3, 4.85)
textbox(s, 7.05, 1.5, 5.8, 0.4, "LOCAL SHAP · one hot, tree-less fringe cell", size=13, color=ACCENT,
        bold=True, spacing=0)
picture(s, local, 7.05, 1.9, 5.7, 3.1)
for i, (big, lab) in enumerate([(f"{N['spectral']} °C", "spectral group, mean |SHAP|"),
                                (f"ρ {N['spearman']}", "same ranking in XGBoost"),
                                (f"{N['group_stab']}%", "same leading group")]):
    rect(s, 7.05 + i * 1.9, 5.25, 1.75, 1.5, TINT)
    callout(s, 7.15 + i * 1.9, 5.3, 1.6, big, lab, big_size=22, label_size=13)

# ---------------------------------------------------------------------------
# 12 Counterfactual analysis
# ---------------------------------------------------------------------------
s = new_slide("Counterfactual analysis", "What if tree cover increases?", f"""
Now we use the trained models as a policy simulator. A counterfactual asks: what would the model predict if
one thing about a place were different? Here the thing we change is tree cover.

Step one, the current state. We selected the urban cells, more than half built-up, that fall in the hottest
ten percent of observed temperature, averaged over three years. That gives {N['targets']} cells. They are
mostly tree-sparse, newer developments on the city fringe, many on the eastern edge, not the old core.

Step two, the modification. We raise their tree cover by ten percentage points. Because fractions must still
add up to one, the new canopy is taken from built-up, bare and grass in proportion to what each cell has.
We do not touch the spectral indices in the lower bound. In the realistic bound, we let NDVI, NDBI and MNDWI
respond, using small models that learn how the indices change with land cover. The truth is likely
between the two, so we always report a range.

Step three, the predicted cooling. Averaged over LightGBM, XGBoost and Random Forest, ten more points of tree
cover go with {N['A_lo']} to {N['A_hi']} °C cooler surfaces. Twenty points go with {N['B_lo']} to
{N['B_hi']} °C. Raising every cell just to the {N['knot']}% knot goes with {N['C_lo']} to {N['C_hi']} °C.
The maps show where the cooling falls for each scenario: blue is cooler.

Two subtleties. First, diminishing returns: the cooling per point of canopy roughly halves beyond about ten to
twelve percent cover, from {N['per_pt_below']} to {N['per_pt_above']} °C per point. So spreading some canopy
across many bare places beats concentrating it in a few. Second, we exclude the very first pixel of canopy.
Going from zero to one percent is worth {N['first_pt']} °C in the model, but that reflects what kind of
place has no trees at all, not what one tree does.

Finally, the caveat that matters for policy. The model learned from comparing places, so these numbers say
how much cooler similar places with more trees are. They are associations, not measured effects of
planting, and they give planners a ranking and a plausible range, not a guarantee.
""")
flow = [("Current", f"{N['targets']} hottest urban cells, tree-sparse", TINT, INK, INK2),
        ("Modified", "+10 points tree cover, taken from built, bare, grass", TINT, INK, INK2),
        ("Predicted cooling", f"{N['A_lo']}–{N['A_hi']} °C cooler (lower–realistic)", ACCENT, WHITE, WHITE)]
for i, (a, b, fill, hc, bc) in enumerate(flow):
    x0 = 0.6 + i * 4.18
    rect(s, x0, 1.55, 3.74, 1.5, fill)
    badge(s, x0 + 0.2, 1.72, 0.5, str(i + 1), fill=WHITE if fill == ACCENT else ACCENT,
          color=ACCENT if fill == ACCENT else WHITE, size=15)
    textbox(s, x0 + 0.85, 1.72, 2.8, 0.5, a, size=18, color=hc, bold=True, spacing=0)
    textbox(s, x0 + 0.2, 2.3, 3.4, 0.75, b, size=15, color=bc, spacing=0)
    if i < 2:
        arrow(s, x0 + 3.8, 2.12)
picture(s, Image.open(FIG / "07_delta_lst_maps.png").crop((380, 80, 4494, 1200)), 0.4, 3.2, 12.5, 3.05)
textbox(s, 0.6, 6.4, 12.1, 0.5, [[("Policy simulation:  ", {"color": ACCENT, "bold": True}),
                                 (f"+20 points → {N['B_lo']}–{N['B_hi']} °C · gains per point halve past "
                                  f"~{N['knot']}% · associations, not causal effects", {})]], size=15)

# ---------------------------------------------------------------------------
# 13 Applications
# ---------------------------------------------------------------------------
s = new_slide("Applications", "From a heat map to decisions", f"""
What can a city actually do with this? Let me go through six applications, each tied to something the
project produced.

Smart cities. The whole pipeline runs on open satellite data and rebuilds every table and figure from one
command. A city could rerun it every pre-monsoon season and get an updated, hundred-metre heat map at
almost no cost, without installing sensors. That is a cheap monitoring layer for a smart-city platform.

Urban planning. The first lesson is where the problem is. In daytime Pune, the hottest urban surfaces are
not in the dense old core; they are on the dry, tree-sparse fringe, where development is newest and still
underway. That is exactly where planning rules for new layouts, setbacks and open spaces can still change
the outcome.

Green infrastructure. Our scenarios suggest that the biggest gains per unit of canopy come from getting
bare places up to about ten percent tree cover. That favours spreading street trees, small parks and green
buffers widely across hot areas, rather than one large park in one place.

Tree plantation. We can rank individual hundred-metre cells: the {N['targets']} hottest urban cells are a
ready-made, mappable priority list for plantation drives, ward by ward.

Heat mitigation. Heat action plans need to know where people are exposed to the hottest surfaces. The map,
together with the per-cell SHAP explanation, tells officials not only where it is hot but why: dryness,
lack of vegetation, or terrain. Different causes call for different fixes, such as greening, irrigation
or cool roofs.

Climate adaptation. Heatwaves are expected to become more frequent. A reproducible, honest model gives a
baseline against which future seasons can be compared, and the same method can be transferred to other
Indian cities with only the study-area box changed.

We also built a read-only Streamlit dashboard that shows every result and lets a user pick a cell, see its
explanation, and try a greening what-if. Throughout, the outputs support decisions; they do not replace
field surveys or local knowledge.
""")
apps = [("Smart cities", "seasonal 100 m heat map from open data"),
        ("Urban planning", "target the dry fringe, not the core"),
        ("Green infrastructure", "spread canopy to ~10% cover"),
        ("Tree plantation", f"priority list: {N['targets']} hottest urban cells"),
        ("Heat mitigation", "know where and why it is hot"),
        ("Climate adaptation", "reproducible baseline for future seasons")]
for i, (a, b) in enumerate(apps):
    row, col = divmod(i, 3)
    x0, y0 = 0.6 + col * 4.12, 1.6 + row * 2.2
    rect(s, x0, y0, 3.86, 1.95, TINT)
    badge(s, x0 + 0.25, y0 + 0.22, 0.6, str(i + 1), size=17)
    textbox(s, x0 + 1.0, y0 + 0.22, 2.8, 0.6, a, size=19, color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE,
            spacing=0)
    textbox(s, x0 + 0.25, y0 + 1.0, 3.4, 0.85, b, size=16, spacing=0)
rect(s, 0.6, 6.0, 12.1, 0.78, NAVY)
textbox(s, 0.85, 6.0, 11.6, 0.78, [[("Dashboard:  ", {"color": WHITE, "bold": True}),
                                   ("pick a cell, read its SHAP breakdown, try a greening what-if",
                                    {"color": PALE})]], size=17, anchor=MSO_ANCHOR.MIDDLE, spacing=0)

# ---------------------------------------------------------------------------
# 14 Limitations
# ---------------------------------------------------------------------------
s = new_slide("Limitations", "What these results cannot tell us", f"""
Every study has limits, and stating them clearly is part of doing it rigorously. Here are five.

First, Pune only. We studied one city with one climate, a semi-arid, hilly setting where the surroundings
dry out before the monsoon. That dryness is exactly why the core is cooler by day. In a humid or irrigated
region the pattern could be different, so we cannot assume the drivers transfer without testing.

Second, pre-monsoon data only, and only daytime. Landsat passes over in the late morning. At night, built
surfaces release stored heat and the classic heat island may well reappear. After the monsoon, the
surroundings turn green and the contrast changes. Our conclusions are for daytime, March to May.

Third, the hundred-metre thermal resolution. Landsat's thermal sensor sees about a hundred metres, so a cell
mixes roofs, roads, trees and courtyards. We cannot resolve a single street or building, and heat blurs
across cell edges. You can see a symptom on the right: this is the map of model residuals, red where a cell
is hotter than predicted, blue where it is cooler. The errors are clustered, with residual Moran's I around
{N['moran']}. That tells us there are drivers we cannot see, such as building height, roof material and
irrigation.

Fourth, correlation is not causation. Everything we learned comes from comparing different places at the
same time. Places with more trees are cooler, but they may also be wealthier, better irrigated or older.
SHAP explains the model, not physics, and the scenarios are associations. Proving a causal effect would need
before-and-after data from actual planting.

Fifth, static land cover. We used one WorldCover map from 2021 for all three years. Any construction or
clearing since then is invisible to the land-cover fractions, though the Sentinel-2 indices, which are
updated every year, partly compensate.

None of these limits changes the validation results, which are honest by design. They bound how far the
scenarios can be read, and they point directly to the future work on the next slide.
""")
lims = [("Pune only", "one city, one climate"),
        ("Pre-monsoon data", "daytime, Mar–May"),
        ("100 m thermal resolution", "mixed pixels"),
        ("Correlation ≠ causation", "associations only"),
        ("Static land cover", "WorldCover 2021")]
for i, (a, b) in enumerate(lims):
    y0 = 1.6 + i * 0.98
    badge(s, 0.6, y0 + 0.08, 0.55, str(i + 1), size=15)
    textbox(s, 1.35, y0, 4.9, 0.85, [[(a, {"bold": True, "color": INK})], [(b, {"size": 15})]], size=18,
            spacing=0)
picture(s, "05_residual_map.png", 6.4, 1.45, 6.6, 5.4)

# ---------------------------------------------------------------------------
# 15 Future work
# ---------------------------------------------------------------------------
s = new_slide("Future work", "Where the work goes next", f"""
The limitations point straight to the next steps. We have arranged them on a timeline, from what could be
done next semester to longer-term research.

In the near term, more cities. The pipeline takes only a study-area box, so the first step is to run it on
other Indian cities with different climates, such as a humid coastal city and a hotter inland one. If the
same drivers come out on top, the conclusions generalise; if not, we learn where they break.

Also near term, multiple seasons. Adding night-time thermal data and post-monsoon and winter composites
would show whether Pune's daytime cool island flips into a classic heat island at night, which matters for
health, because night-time heat prevents the body from recovering.

In the medium term, higher-resolution imagery. Thermal data from the ECOSTRESS instrument on the space
station, at roughly seventy metres, or downscaling methods, would sharpen the picture. Building footprints
and heights would add the urban-form drivers that our clustered residuals point to.

Also medium term, meteorological variables. Air temperature, humidity and wind from weather stations or
reanalysis data would let us connect surface temperature to the air temperature people actually feel, and
test whether our associations carry over.

In the long term, a real-time dashboard. Our current Streamlit dashboard is read-only and shows stored
results. The next version would ingest each new season automatically, update the map, and let planners run
the realistic-bound scenario live.

And finally, climate projections. Combining our model with downscaled climate-model scenarios would let us
ask how much canopy would be needed to offset projected warming in the hottest wards by mid-century. That
would turn the tool from describing today into planning for tomorrow.

On the method side, we would also add nested spatial cross-validation to remove the small optimism left
by tuning, and before-and-after comparisons of land-cover change to move from association toward cause.
""")
fut = [("More cities", "test whether drivers transfer"), ("Multiple seasons", "night, post-monsoon, winter"),
       ("Higher resolution", "~70 m thermal, building heights"), ("Meteorology", "air temperature, humidity, wind"),
       ("Real-time dashboard", "auto-update every season"), ("Climate projections", "canopy needed for future warming")]
line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.2), Inches(3.05), Inches(12.1), Inches(3.05))
line.line.color.rgb = ACCENT
line.line.width = Pt(3)
for j, hz in enumerate(["Near term", "Medium term", "Long term"]):
    textbox(s, 0.6 + j * 4.1, 1.7, 3.9, 0.45, hz, size=18, color=ACCENT, bold=True, align=PP_ALIGN.CENTER,
            spacing=0)
for i, (a, b) in enumerate(fut):
    cx = 1.6 + i * 2.05
    badge(s, cx - 0.3, 2.75, 0.6, str(i + 1), size=16)
    rect(s, cx - 0.95, 3.75, 1.9, 2.0, TINT)
    textbox(s, cx - 0.85, 3.9, 1.7, 0.75, a, size=17, color=INK, bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE, spacing=0)
    textbox(s, cx - 0.85, 4.75, 1.7, 0.95, b, size=14, align=PP_ALIGN.CENTER, spacing=0)
textbox(s, 0.6, 6.05, 12.1, 0.6, [[("Method:  ", {"color": ACCENT, "bold": True}),
                                  ("nested spatial CV · before-and-after land-cover change, toward causal "
                                   "evidence", {})]], size=17)

# ---------------------------------------------------------------------------
# 16 Conclusion
# ---------------------------------------------------------------------------
s = new_slide("Conclusion", "From satellite pixels to planning evidence", f"""
Let me pull everything together, following the flow on the slide.

The problem. Pune is urbanising fast, and heat action plans need to know where greening will cool the most.
The textbook assumption of a hot core does not hold in daytime Pune: the core is cooler than the dry
fringe. And common evaluation practice, random pixel splits, overstates how well models work.

The method. We combined four open satellite datasets into {N['n_features']} physically meaningful features
on a hundred-metre grid, modelled the temperature anomaly with five models, and evaluated them with
{N['block_km']}-kilometre spatial block cross-validation and an untouched 2025 test, never using coordinates
as features.

The results. Random cross-validation flattered the flexible models by up to {N['gap_tree_hi']} °C. Under honest
testing, the tree models tie, and LightGBM, our final model, predicts the anomaly in unseen places and an
unseen year with an RMSE of {N['rmse_2025']} °C and R squared of {N['r2_2025']}.

The explainability. SHAP shows that vegetation, dryness and moisture lead, and that the built-up fraction,
the concrete in our title, is the smallest driver group. The raw link between built-up cover and cooler
surfaces, {N['rq5_raw']} °C per ten points, shrinks by {N['rq5_shrink']}% once vegetation and dryness are
modelled. So the cool core is mainly a vegetation and dryness contrast. These rankings hold across models.

The real-world impact. Pune's hottest urban surfaces are tree-sparse fringe developments. In these data,
raising them to about ten percent tree cover is associated with roughly {N['A_lo1']} to {N['A_hi1']} °C cooler
daytime surfaces, with gains per point roughly halving beyond that. That gives planners a concrete,
mappable target, with honest bounds and clearly stated caveats.

In one sentence: in daytime Pune, the surface heat is in the dry fringe, not the dense core, and canopy
matters far more than concrete.

Thank you for your attention. We are happy to take your questions, and every number you have seen can be
traced back to an output table in our repository.
""")
chev = [("Problem", "where is Pune hot, and why?"),
        ("Method", f"{N['n_features']} features · {N['block_km']} km spatial CV"),
        ("Results", f"{N['rmse_2025']} °C RMSE · R² {N['r2_2025']} on 2025"),
        ("Explainability", "vegetation and dryness lead; built-up smallest"),
        ("Impact", f"~10% canopy: {N['A_lo1']}–{N['A_hi1']} °C cooler")]
for i, (a, b) in enumerate(chev):
    x0 = 0.6 + i * 2.4
    shp = s.shapes.add_shape(MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, Inches(x0), Inches(1.6),
                             Inches(2.42), Inches(0.9))
    shp.fill.solid()
    shp.fill.fore_color.rgb = NAVY if i == 4 else ACCENT
    shp.line.fill.background()
    shp.shadow.inherit = False
    textbox(s, x0 + (0.1 if i == 0 else 0.4), 1.6, 1.9, 0.9, a, size=16, color=WHITE, bold=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=0)
    textbox(s, x0 + 0.15, 2.7, 2.2, 1.3, b, size=15, color=INK2, spacing=0)
rect(s, 0.6, 4.0, 12.1, 2.75, NAVY)
textbox(s, 1.0, 4.3, 6.0, 1.0, "Thank you", size=44, color=WHITE, bold=True, spacing=0)
textbox(s, 1.0, 5.4, 6.0, 0.8, "Questions?", size=32, color=ACCENT, bold=True, spacing=0)
textbox(s, 7.2, 4.5, 5.3, 2.0, [" · ".join(TEAM[:2]), " · ".join(TEAM[2:]),
                                [("Symbiosis Institute of Technology, Pune", {"color": PALE, "size": 14})]],
        size=17, color=WHITE, spacing=6)

# ---------------------------------------------------------------------------
# Checks and save
# ---------------------------------------------------------------------------
assert len(prs.slides) == TOTAL
prs.save(OUT)
sys.stdout.reconfigure(encoding="utf-8")
print(f"wrote {OUT.relative_to(ROOT)}")
WPM = 140
for i, sl in enumerate(prs.slides, 1):
    words = len(sl.notes_slide.notes_text_frame.text.split())
    print(f"slide {i:2d}: {words} note words ≈ {words / WPM:.1f} min")
total = sum(len(sl.notes_slide.notes_text_frame.text.split()) for sl in prs.slides)
print(f"total {total} words ≈ {total / WPM:.0f} min at {WPM} wpm")
