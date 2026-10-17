"""T8.2 / T8.4 / T8.5: check that quoted numbers match the output tables, in the
report (reports/report.md), the slide deck (reports/MirageAI_Slides.pptx, slide
text and speaker notes) and the viva notes (reports/viva_prep.md).

    python scripts/check_report_numbers.py

Each check recomputes a value from outputs/tables (or, where a number was
derived in a phase findings file, from that file), formats it at the precision
the document uses, and requires it to appear in a paragraph (report, viva) or
on a slide / in its notes (deck) that also contains the given keyword. A check
fails if the number is missing or different, so a stale or mistyped figure
cannot survive a table update. For the deck and the viva notes, every number
in the text must also be covered: a number no check verified, and that is not
a design constant confirmed against config.py, is flagged too. Documents that
do not exist yet are skipped. Exit code 1 if anything fails.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pandas as pd

from mirageai.results import load_best_params

ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / "outputs" / "tables"
FINDINGS = {p: (ROOT / "reports" / f"{p}_findings.md").read_text(encoding="utf-8")
            for p in ("phase6", "phase7")}


def _md_paras(path: Path) -> list[str]:
    if not path.exists():
        return []
    return [re.sub(r"\s+", " ", p) for p in re.split(r"\n\s*\n", path.read_text(encoding="utf-8"))]


def _deck_paras(path: Path) -> list[str]:
    """One paragraph per slide (all shapes, tables included) and one per notes page."""
    if not path.exists():
        return []
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        parts = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts += [cell.text for row in sh.table.rows for cell in row.cells]
        # The footer and "n / N" slide number are layout, not quoted numbers.
        body = re.sub(r"(?<![\d.])\d+ / \d+(?![\d.])|MirageAI · Concrete vs\. Canopy", " ", " | ".join(parts))
        out.append(re.sub(r"\s+", " ", f"[slide {i}] " + body))
        out.append(re.sub(r"\s+", " ", f"[notes {i}] " + slide.notes_slide.notes_text_frame.text))
    return out


DOCS = {
    "report": _md_paras(ROOT / "reports" / "report.md"),
    "slides": _deck_paras(ROOT / "reports" / "MirageAI_Slides.pptx"),
    "viva": _md_paras(ROOT / "reports" / "viva_prep.md"),
}
PARAS = DOCS["report"]
COVERED: dict[str, set[str]] = {d: set() for d in DOCS}


def T(name: str) -> pd.DataFrame:
    return pd.read_csv(TABLES / f"{name}.csv")


def fmt(v: float, nd: int, thousands: bool = False) -> str:
    s = f"{abs(v):,.{nd}f}" if thousands else f"{abs(v):.{nd}f}"
    return s


results = []


def check(label: str, keyword: str, value: float, nd: int, thousands: bool = False,
          doc: str = "report") -> None:
    if not DOCS[doc]:
        return
    s = fmt(value, nd, thousands)
    paras = [p for p in DOCS[doc] if keyword.lower() in p.lower()]
    pat = re.compile(r"(?<![\d.,])" + re.escape(s) + r"(?![\d])")
    ok = any(pat.search(p) for p in paras)
    if ok:
        COVERED[doc].add(s)
    results.append((ok, f"[{doc}] {label}", s, keyword, len(paras)))


def check_findings(label: str, keyword: str, text: str, phase: str, doc: str = "report") -> None:
    """A number taken from a findings file: it must appear in both."""
    if not DOCS[doc]:
        return
    in_doc = any(keyword.lower() in p.lower() and text in p for p in DOCS[doc])
    in_source = text in FINDINGS[phase]
    if in_doc and in_source:
        COVERED[doc].update(re.findall(r"\d[\d,]*(?:\.\d+)?", text))
    results.append((in_doc and in_source, f"[{doc}] {label}", text, keyword, int(in_source)))


# --- Data and water exclusion (phase4_water_exclusion) ----------------------
w = T("phase4_water_exclusion").set_index("year")
check("cells per year", "81,983 cells", w.loc["2023", "rows_before"], 0, True)
check("cell-years total", "245,949", w.loc["all", "rows_before"], 0, True)
check("water rows removed", "water", w.loc["all", "rows_removed"], 0, True)
check("water cells removed", "water", w.loc["2023", "cells_removed"], 0, True)
check("rows after", "water", w.loc["all", "rows_after"], 0, True)
check("land cells per year", "water", w.loc["2023", "rows_after"], 0, True)
shift = w.loc[["2023", "2024", "2025"], "median_LST_shift_if_land_only"].astype(float)
check("median shift min", "land-only median", shift.min(), 3)
check("median shift max", "land-only median", shift.max(), 3)

# --- Cool island, land cells (phase8_core_ring_land) -----------------------
p1 = T("phase8_core_ring_land").set_index("year")
for y in (2023, 2024, 2025):
    check(f"core minus ring {y}", "outer ring", p1.loc[y, "core_minus_ring"], 3)
check("core NDVI 2024", "outer ring", p1.loc[2024, "core_NDVI"], 3)
check("ring NDVI 2024", "outer ring", p1.loc[2024, "ring_NDVI"], 3)

# --- Residual range (phase4_residual_range, phase8_residual_range_by_scheme) -
rr = T("phase4_residual_range").set_index("year")
for y in (2023, 2024, 2025):
    check(f"gate range {y}", "residual variogram", rr.loc[y, "range_95_m"], 0, True)
check("gate median", "residual variogram", rr.range_95_m.median(), 0, True)
rs = T("phase8_residual_range_by_scheme")
check("tuned 4 km median", "tuned model's own residuals",
      rs[(rs.scheme == "spatial_4km") & (rs.seed == 42)].median_range_95_m.iloc[0], 0, True)
check("random-CV range", "residual variogram",
      rs[rs.scheme == "random"].median_range_95_m.iloc[0], 0, True)
check("8 km range", "residual variogram", rs[rs.scheme == "spatial_8km"].median_range_95_m.iloc[0], 0, True)

# --- Model comparison (model_comparison) ------------------------------------
mc = T("model_comparison").set_index(["scheme", "model"])
for m in ("baseline", "linear", "rf", "xgb", "lgbm"):
    for s, kw in (("random", "Table 1."), ("spatial_4km", "Table 1."), ("spatio_temporal_4km", "Table 1.")):
        check(f"{s} {m} RMSE", "| " + {"baseline": "Baseline (mean)", "linear": "Linear regression",
                                        "rf": "Random Forest", "xgb": "XGBoost",
                                        "lgbm": "LightGBM"}[m] + " |", mc.loc[(s, m), "RMSE_mean"], 3)
        check(f"{s} {m} RMSE std", "| " + {"baseline": "Baseline (mean)", "linear": "Linear regression",
                                            "rf": "Random Forest", "xgb": "XGBoost",
                                            "lgbm": "LightGBM"}[m] + " |", mc.loc[(s, m), "RMSE_std"], 3)
check("lgbm 4 km in text", "tied", mc.loc[("spatial_4km", "lgbm"), "RMSE_mean"], 3)
check("xgb random in text", "XGBoost is clearly", mc.loc[("random", "xgb"), "RMSE_mean"], 3)
for s, label in (("random", "random"), ("spatial_1km", "1 km"), ("spatial_2km", "2 km"),
                 ("spatial_8km", "8 km")):
    check(f"lgbm curve {label}", "rises steadily", mc.loc[(s, "lgbm"), "RMSE_mean"], 3)
check("lgbm R2 random", "rises steadily", mc.loc[("random", "lgbm"), "R2_mean"], 3)
check("lgbm R2 8 km", "rises steadily", mc.loc[("spatial_8km", "lgbm"), "R2_mean"], 3)
pv = T("phase5_paired_vs_best").set_index(["scheme", "model"])
check("xgb vs lgbm 4 km paired std", "paired", pv.loc[("spatial_4km", "xgb"), "dRMSE_std"], 3)
check("xgb better runs random", "all 25 paired runs", pv.loc[("random", "xgb"), "runs_model_better"], 0)

# --- RQ3 gap (phase5_cv_gap, phase5_strip_stress_test) ----------------------
gap = T("phase5_cv_gap").set_index(["model", "scheme"])
for m in ("rf", "lgbm", "xgb", "linear"):
    check(f"gap {m}", "inflation depends on the model", gap.loc[(m, "spatial_4km"), "dRMSE"], 3)
check("gap pct rf", "inflation depends", gap.loc[("rf", "spatial_4km"), "dRMSE_pct"], 1)
check("gap pct xgb", "inflation depends", gap.loc[("xgb", "spatial_4km"), "dRMSE_pct"], 1)
check("gap pct linear", "inflation depends", gap.loc[("linear", "spatial_4km"), "dRMSE_pct"], 1)
st = T("phase5_strip_stress_test").set_index(["scheme", "model"])
check("strip penalty", "strip layout", st.loc[("spatial_4km", "lgbm"), "dRMSE_strips_minus_shuffled"], 3)
check("strip RMSE", "strip layout", st.loc[("spatial_4km", "lgbm"), "RMSE_strips_mean"], 3)

# --- 2025 test and memorisation (phase5_temporal_holdout) -------------------
th = T("phase5_temporal_holdout").set_index(["scheme", "model"])
k = "spatio_temporal_4km"
check("lgbm unseen", "never seen in training", th.loc[(k, "lgbm"), "RMSE_unseen_mean"], 3)
check("lgbm unseen R2", "never seen in training", th.loc[(k, "lgbm"), "R2_unseen_mean"], 3)
check("lgbm seen", "never seen in training", th.loc[(k, "lgbm"), "RMSE_seen_place_mean"], 3)
check("lgbm seen R2", "never seen in training", th.loc[(k, "lgbm"), "R2_seen_place_mean"], 3)
for m in ("lgbm", "xgb", "rf", "linear"):
    check(f"memorisation {m}", "memorisation gap", th.loc[(k, m), "memorisation_dRMSE_mean"], 3)

# --- Errors (phase5_urban_vs_rest, overlap, Moran) --------------------------
uv = T("phase5_urban_vs_rest").set_index(["scheme", "model", "urban_group"])
check("urban RMSE", "lower R²", uv.loc[("spatial_4km", "lgbm", "urban"), "RMSE_mean"], 3)
check("rest RMSE", "lower R²", uv.loc[("spatial_4km", "lgbm", "rest"), "RMSE_mean"], 3)
check("urban R2", "lower R²", uv.loc[("spatial_4km", "lgbm", "urban"), "R2_mean"], 3)
check("rest R2", "lower R²", uv.loc[("spatial_4km", "lgbm", "rest"), "R2_mean"], 3)
check("urban baseline", "lower R²", uv.loc[("spatial_4km", "baseline", "urban"), "RMSE_mean"], 3)
check("rest baseline", "lower R²", uv.loc[("spatial_4km", "baseline", "rest"), "RMSE_mean"], 3)
ov = T("phase5_hot_decile_overlap").groupby(["scheme", "decile", "model"]).overlap.mean()
check("hot AOI", "hottest tenth", 100 * ov.loc[("spatial_4km", "AOI-wide", "lgbm")], 1)
check("hot urban", "hottest tenth", 100 * ov.loc[("spatial_4km", "within urban", "lgbm")], 1)
check("54% precision in methods", "precision", round(100 * ov.loc[("spatial_4km", "within urban", "lgbm")]), 0)
mi = T("phase5_morans_i")
lg = mi[(mi.scheme == "spatial_4km") & (mi.model == "lgbm")].morans_I
raw = mi[mi.model == "raw_target"].morans_I
check("moran lgbm min", "Moran", lg.min(), 3)
check("moran lgbm max", "Moran", lg.max(), 3)
check("moran raw min", "Moran", raw.min(), 3)
check("moran raw max", "Moran", raw.max(), 3)

# --- SHAP (phase6_*) --------------------------------------------------------
imp = T("phase6_shap_importance").set_index("feature")
ROW = {"NDVI": "| NDVI |", "NDBI": "| Dryness index (NDBI) |", "MNDWI": "| Moisture index (MNDWI) |",
       "frac_grass": "| Grass fraction |", "frac_tree": "| Tree cover fraction |",
       "frac_crop": "| Cropland fraction |", "elevation": "| Elevation |",
       "frac_built": "| Built-up fraction |"}
for f, row in ROW.items():
    check(f"imp {f} all", row, imp.loc[f, "lgbm_all"], 3)
    check(f"imp {f} urban", row, imp.loc[f, "lgbm_urban"], 3)
check("NDVI in text", "**NDVI:**", imp.loc["NDVI", "lgbm_all"], 3)
check("slope in text", "terrain", imp.loc["slope", "lgbm_all"], 3)
check("dist_water in text", "terrain", imp.loc["dist_water_m", "lgbm_all"], 3)
gs = T("phase6_grouped_spectral").set_index(["model", "subset"])
check("spectral block", "| Spectral block (sum", gs.loc[("lgbm", "all"), "group_mean_abs"], 3)
check("spectral block urban", "| Spectral block (sum", gs.loc[("lgbm", "urban"), "group_mean_abs"], 3)
check("cancellation", "cancel", 100 * gs.loc[("lgbm", "all"), "cancellation_share"], 0)
check("block / next", "2.75 times",
      gs.loc[("lgbm", "all"), "group_mean_abs"] / gs.loc[("lgbm", "all"), "top_other_mean_abs"], 2)
ra = T("phase6_rank_agreement").set_index("comparison")
check("spearman all", "Spearman", ra.loc["lgbm vs xgb, all cells", "spearman"], 3)
check("spearman urban", "Spearman", ra.loc["lgbm vs xgb, urban", "spearman"], 3)
ds = T("phase6_dominant_stability").set_index("level")
check("group stability", "driver *groups*", 100 * ds.loc["group", "lgbm_vs_xgb"], 1)
check("feature stability", "driver *groups*", 100 * ds.loc["feature", "lgbm_vs_xgb"], 1)
dg = T("phase6_dominant_driver_groups").set_index("dominant")
check("spectral share", "driver *groups*", 100 * dg.loc["spectral", "share"], 1)
check("dry share", "driver *groups*", 100 * dg.loc["dry cover", "share"], 1)
check("dry share urban", "driver *groups*", 100 * dg.loc["dry cover", "share_of_urban"], 1)
inter = T("phase6_interactions")
for i in range(3):
    check(f"interaction {i + 1}", "interactions", inter.mean_abs_interaction.iloc[i], 3)
thr = T("phase6_thresholds").set_index("feature")
t = thr.loc["frac_tree"]
check("tree step", "first canopy", t.step_zero_to_low, 2)
check("tree slope below", "per 10 points", t.pos_slope_below / 10, 2)
check("tree slope above", "per 10 points", t.pos_slope_above / 10, 2)
check("tree knot", "knot at", 100 * t.pos_knot, 1)
check("tree knot lo", "knot at", 100 * t.pos_knot_lo, 1)
check("tree knot hi", "knot at", 100 * t.pos_knot_hi, 1)
check("tree r2 linear", "straight", 100 * t.r2_linear, 0)
check("tree r2 seg", "straight", 100 * t.r2_segmented, 0)
check("MNDWI onset", "no association below", thr.loc["MNDWI", "knot"], 2)
check("MNDWI slope", "no association below", thr.loc["MNDWI", "slope_above"] / 10, 1)
check("NDVI knot", "NDVI is associated", thr.loc["NDVI", "knot"], 2)
check("NDVI slope below", "NDVI is associated", thr.loc["NDVI", "slope_below"] / 10, 1)
check("NDVI slope above", "NDVI is associated", thr.loc["NDVI", "slope_above"] / 10, 1)
check("water reach", "water reaches", round(thr.loc["dist_water_m", "knot"], -1), 0)
check("elevation knot", "elevation", round(thr.loc["elevation", "knot"], -1), 0)
check("linear NDBI coef", "standardised dryness", thr.loc["NDBI", "linear_coef_per_sd"], 2)
check("tree NDBI response", "standardised dryness", thr.loc["NDBI", "tree_linearised_per_sd"], 2)
check("linear NDVI coef", "standardised dryness", thr.loc["NDVI", "linear_coef_per_sd"], 2)
check("tree NDVI response", "standardised dryness", thr.loc["NDVI", "tree_linearised_per_sd"], 2)
fc = T("phase8_feature_correlations").set_index("feature")
check("NDBI raw r", "standardised dryness", fc.loc["NDBI", "r_with_LST_anom"], 2)
check("NDBI raw r discussion", "not a built-up index", fc.loc["NDBI", "r_with_LST_anom"], 2)
ic = T("phase8_index_correlations")
check("NDVI-NDBI r", "strongly correlated", ic[(ic.a == "NDVI") & (ic.b == "NDBI")].r.iloc[0], 2)
check_findings("block total r", "spectral-block total", "r = 0.96", "phase6")
check_findings("linear RMSE cost", "spatial-CV RMSE", "0.13 °C", "phase6")
rq5 = T("phase6_rq5_frac_built").set_index("subset")
check("frac_built raw r", "cool-island association", rq5.loc["all", "raw_pearson"], 2)
check("frac_built raw slope", "cool-island association", rq5.loc["all", "raw_slope_per_0.1"], 2)
check("frac_built shap slope", "cool-island association", rq5.loc["all", "shap_slope_per_0.1"], 3)
check("frac_built shrink", "cool-island association",
      100 * (1 - rq5.loc["all", "shap_slope_per_0.1"] / rq5.loc["all", "raw_slope_per_0.1"]), 0)
check("urban raw r", "cool-island association", rq5.loc["urban", "raw_pearson"], 3)

# --- Scenarios (phase7_*) ---------------------------------------------------
tc = T("phase7_target_cells")
check("targets", "targets are", len(tc), 0, True)
cv7 = T("phase7_index_model_cv").groupby("index").R2.mean()
for idx in ("MNDWI", "NDVI", "NDBI"):
    check(f"index R2 {idx}", "index models", cv7[idx], 3)
sm = T("phase7_summary")
ens = sm[sm.model == "ensemble"].set_index(["scenario", "bound"])
ROW3 = {"A": "| A: +10 points |", "B": "| B: +20 points |", "C": "| C: to the 10.2% knot"}
for s, row in ROW3.items():
    for b in ("lower", "realistic"):
        check(f"{s} {b} headline", row, ens.loc[(s, b), "dLST_excl_first_point"], 2)
        check(f"{s} {b} min", row, ens.loc[(s, b), "dLST_excl_min"], 2)
        check(f"{s} {b} max", row, ens.loc[(s, b), "dLST_excl_max"], 2)
check("C points", ROW3["C"], ens.loc[("C", "lower"), "mean_points_added"], 1)
check("headline bracket A text", "Raising the targets", ens.loc[("A", "lower"), "dLST_excl_first_point"], 2)
check("headline bracket B text", "Doubling", ens.loc[("B", "lower"), "dLST_excl_first_point"], 2)
check("max flagged", "training data", ens.pct_flagged.max(), 2)
dr = T("phase7_dose_response")
check("flagged at 30", "training data", dr[(dr.points == 30) & (dr.bound == "lower")].pct_flagged.max(), 2)
ms = T("phase7_marginal_summary").set_index("bound")
check("first point realistic", "first point of canopy", ms.loc["realistic", "first_point_0_to_1pct"], 2)
check("first point lower", "first point of canopy", ms.loc["lower", "first_point_0_to_1pct"], 2)
check("per point below knot", "roughly halve", ms.loc["realistic", "per_point_2pct_to_knot"], 2)
check("per point above", "roughly halve", ms.loc["realistic", "per_point_12_to_30pct"], 2)
check_findings("targets with no tree cover", "no tree cover", "78.7%", "phase7")
check_findings("target threshold", "+2.41", "+2.41 °C", "phase7")
check_findings("NE share", "north-east", "45%", "phase7")
check_findings("SE share", "south-east", "29%", "phase7")
check_findings("median distance", "median 11.5 km", "11.5 km", "phase7")
check_findings("built median", "median built-up", "78%", "phase7")
check_findings("urban cells", "2,673 of 26,726", "26,726", "phase7")

# --- Hyperparameters (best_params.json, or the shipped tuning tables without it) ---
best = load_best_params()
for m, label in (("lgbm", "| LightGBM | n_estimators"), ("xgb", "| XGBoost | n_estimators"),
                 ("rf", "| Random Forest | n_estimators")):
    check(f"tuning RMSE {m}", label, best[m]["best_cv_rmse"], 4)
    for p, v in best[m]["best_params"].items():
        results.append((any(label in q and f"{p} {v}" in q for q in PARAS),
                        f"param {m}.{p}", f"{p} {v}", label, 1))

# ===========================================================================
# Slide deck (doc="slides"): "[slide n]" is the slide, "[notes n]" its notes
# ===========================================================================
S = "slides"
p1 = T("phase8_core_ring_land")
gi = T("phase6_group_importance").set_index("group")
cib = T("phase8_cool_island_bins")
_man = ROOT / "reports" / "slides_manifest.csv"
# Slide numbers by title, from the manifest the deck build writes, so a reordered deck keeps its checks.
SN = dict(zip(*pd.read_csv(_man)[["title", "slide"]].T.values)) if DOCS[S] and _man.exists() else {}


def both(label: str, title: str, value: float, nd: int, thousands: bool = False, where=("slide", "notes")):
    if not SN:
        return
    n = SN[title]
    for w_ in where:
        check(f"{label} ({w_} {n})", f"[{w_} {n}]", value, nd, thousands, doc=S)


both("core cooler, low", "Motivation", p1.core_minus_ring.max(), 1)
both("core cooler, high", "Motivation", p1.core_minus_ring.min(), 1)
both("land cell-years", "Dataset", w.loc["all", "rows_after"], 0, True)
both("land cells per year", "Dataset", w.loc["2023", "rows_after"], 0, True, where=("notes",))
both("water rows", "Dataset", w.loc["all", "rows_removed"], 0, True, where=("notes",))
both("residual range", "Methodology", rr.range_95_m.median(), 0, True)
lo_bin = cib[cib.mid < 10].groupby("year").anom.mean()
hi_bin = cib[(cib.mid > 80) & (cib.mid < 90)].groupby("year").anom.mean()
both("built gradient", "Exploratory data analysis", (lo_bin - hi_bin).mean(), 1)
both("gap linear", "Evaluation", gap.loc[("linear", "spatial_4km"), "dRMSE"], 2)
tree_gaps = [gap.loc[(m, "spatial_4km"), "dRMSE"] for m in ("rf", "xgb", "lgbm")]
both("gap trees low", "Evaluation", min(tree_gaps), 2)
both("gap trees high", "Evaluation", max(tree_gaps), 2)
both("gap xgb", "Evaluation", gap.loc[("xgb", "spatial_4km"), "dRMSE"], 2, where=("notes",))
for m in ("linear", "rf", "xgb", "lgbm"):
    both(f"table random {m}", "Results: accuracy", mc.loc[("random", m), "RMSE_mean"], 2, where=("slide",))
    both(f"table 4 km {m}", "Results: accuracy", mc.loc[("spatial_4km", m), "RMSE_mean"], 2, where=("slide",))
    both(f"table unseen {m}", "Results: accuracy", th.loc[(k, m), "RMSE_unseen_mean"], 2, where=("slide",))
    both(f"table seen {m}", "Results: accuracy", th.loc[(k, m), "RMSE_seen_place_mean"], 2, where=("slide",))
tree_4km = [mc.loc[("spatial_4km", m), "RMSE_mean"] for m in ("rf", "xgb", "lgbm")]
both("tied low", "Results: accuracy", min(tree_4km), 2)
both("tied high", "Results: accuracy", max(tree_4km), 2)
both("R2 2025", "Results: accuracy", th.loc[(k, "lgbm"), "R2_unseen_mean"], 2)
both("RMSE 2025", "Results: accuracy", th.loc[(k, "lgbm"), "RMSE_unseen_mean"], 2, where=("notes",))
mem = [th.loc[(k, m), "memorisation_dRMSE_mean"] for m in ("rf", "xgb", "lgbm")]
both("memorisation low", "Results: accuracy", min(mem), 2)
both("memorisation high", "Results: accuracy", max(mem), 2)
both("memorisation linear", "Results: accuracy", th.loc[(k, "linear"), "memorisation_dRMSE_mean"], 2)
both("spectral block", "Results: drivers", gi.loc["spectral", "lgbm_all"], 2)
both("spectral / dry", "Results: drivers", gi.loc["spectral", "lgbm_all"] / gi.loc["dry cover", "lgbm_all"], 1)
both("dry cover", "Results: drivers", gi.loc["dry cover", "lgbm_all"], 2, where=("notes",))
both("dry cover urban", "Results: drivers", gi.loc["dry cover", "lgbm_urban"], 2, where=("notes",))
both("spearman", "Results: drivers", ra.loc["lgbm vs xgb, all cells", "spearman"], 2, where=("notes",))
both("dominant spectral", "Results: drivers", 100 * dg.loc["spectral", "share"], 1, where=("notes",))
both("dominant dry", "Results: drivers", 100 * dg.loc["dry cover", "share"], 1, where=("notes",))
both("dominant dry urban", "Results: drivers", 100 * dg.loc["dry cover", "share_of_urban"], 1, where=("notes",))
both("group stability", "Results: drivers", 100 * ds.loc["group", "lgbm_vs_xgb"], 0)
both("targets", "Results: greening", len(tc), 0, True, where=("notes",))
for s_, b_ in (("A", "lower"), ("A", "realistic"), ("B", "lower"), ("B", "realistic")):
    both(f"{s_} {b_}", "Results: greening", ens.loc[(s_, b_), "dLST_excl_first_point"], 2, where=("notes",))
both("A lower 1dp", "Results: greening", ens.loc[("A", "lower"), "dLST_excl_first_point"], 1, where=("slide",))
both("A realistic 1dp", "Results: greening", ens.loc[("A", "realistic"), "dLST_excl_first_point"], 1,
     where=("slide",))
both("per point to knot", "Results: greening", ms.loc["realistic", "per_point_2pct_to_knot"], 2, where=("notes",))
both("per point past knot", "Results: greening", ms.loc["realistic", "per_point_12_to_30pct"], 2,
     where=("notes",))
both("first point", "Results: greening", ms.loc["realistic", "first_point_0_to_1pct"], 1, where=("notes",))
both("knot", "Key insights", 100 * t.pos_knot, 1)
both("slope below", "Key insights", t.pos_slope_below / 10, 2, where=("notes",))
both("slope above", "Key insights", t.pos_slope_above / 10, 2, where=("notes",))
both("RQ5 raw", "Key insights", rq5.loc["all", "raw_slope_per_0.1"], 2, where=("notes",))
both("RQ5 SHAP", "Key insights", rq5.loc["all", "shap_slope_per_0.1"], 3, where=("notes",))
both("RQ5 shrink", "Key insights",
     100 * (1 - rq5.loc["all", "shap_slope_per_0.1"] / rq5.loc["all", "raw_slope_per_0.1"]), 0)
both("Moran", "Limitations", lg.mean(), 1)
both("conclusion RMSE", "Conclusion", th.loc[(k, "lgbm"), "RMSE_unseen_mean"], 2)
both("conclusion R2", "Conclusion", th.loc[(k, "lgbm"), "R2_unseen_mean"], 2)
both("planner low", "Conclusion", ens.loc[("A", "lower"), "dLST_excl_first_point"], 1)
both("planner high", "Conclusion", ens.loc[("A", "realistic"), "dLST_excl_first_point"], 1)

# ===========================================================================
# Viva notes (doc="viva"): each cheat-sheet row by its label
# ===========================================================================
V = "viva"
imp6 = T("phase6_shap_importance").set_index("feature")
gs6 = T("phase6_grouped_spectral").set_index(["model", "subset"])
ic8 = T("phase8_index_correlations")
cv7 = T("phase7_index_model_cv").groupby("index").R2.mean()
rs8 = T("phase8_residual_range_by_scheme")
st5 = T("phase5_strip_stress_test").set_index(["scheme", "model"])
ov5 = T("phase5_hot_decile_overlap").groupby(["scheme", "decile", "model"]).overlap.mean()
uv5 = T("phase5_urban_vs_rest").set_index(["scheme", "model", "urban_group"])
VIVA = [
    ("Land cells per year", w.loc["2023", "rows_after"], 0, True),
    ("Land cell-years", w.loc["all", "rows_after"], 0, True),
    ("Water-majority rows removed", w.loc["all", "rows_removed"], 0, True),
    ("Core cooler than outer ring", p1.core_minus_ring.max(), 1, False),
    ("Core cooler than outer ring", p1.core_minus_ring.min(), 1, False),
    ("pre-registered gate", rr.range_95_m.median(), 0, True),
    ("tuned model at 4 km", rs8[(rs8.scheme == "spatial_4km") & (rs8.seed == 42)].median_range_95_m.iloc[0], 0, True),
    ("random CV vs 8 km CV", rs8[rs8.scheme == "random"].median_range_95_m.iloc[0], 0, True),
    ("random CV vs 8 km CV", rs8[rs8.scheme == "spatial_8km"].median_range_95_m.iloc[0], 0, True),
    *[("Spatial 4 km RMSE", mc.loc[("spatial_4km", m_), "RMSE_mean"], 3, False) for m_ in ("lgbm", "xgb", "rf", "linear")],
    ("Random CV RMSE, XGBoost", mc.loc[("random", "xgb"), "RMSE_mean"], 3, False),
    ("2025 unseen-place RMSE", th.loc[(k, "lgbm"), "RMSE_unseen_mean"], 3, False),
    ("2025 unseen-place RMSE", th.loc[(k, "lgbm"), "R2_unseen_mean"], 3, False),
    ("2025 seen-place RMSE", th.loc[(k, "lgbm"), "RMSE_seen_place_mean"], 3, False),
    *[("Memorisation gap", th.loc[(k, m_), "memorisation_dRMSE_mean"], 3, False) for m_ in ("lgbm", "xgb", "rf", "linear")],
    ("inflation: linear vs trees", gap.loc[("linear", "spatial_4km"), "dRMSE"], 3, False),
    ("inflation: linear vs trees", min(tree_gaps), 3, False),
    ("inflation: linear vs trees", max(tree_gaps), 3, False),
    ("Strip-layout stress penalty", st5.loc[("spatial_4km", "lgbm"), "dRMSE_strips_minus_shuffled"], 3, False),
    ("Urban vs rest RMSE", uv5.loc[("spatial_4km", "lgbm", "urban"), "RMSE_mean"], 3, False),
    ("Urban vs rest RMSE", uv5.loc[("spatial_4km", "lgbm", "rest"), "RMSE_mean"], 3, False),
    ("Urban vs rest R²", uv5.loc[("spatial_4km", "lgbm", "urban"), "R2_mean"], 3, False),
    ("Urban vs rest R²", uv5.loc[("spatial_4km", "lgbm", "rest"), "R2_mean"], 3, False),
    ("Hottest-decile detection", 100 * ov5.loc[("spatial_4km", "within urban", "lgbm")], 1, False),
    ("Hottest-decile detection", 100 * ov5.loc[("spatial_4km", "AOI-wide", "lgbm")], 1, False),
    ("Residual Moran's I", lg.min(), 3, False), ("Residual Moran's I", lg.max(), 3, False),
    ("Residual Moran's I", raw.min(), 3, False), ("Residual Moran's I", raw.max(), 3, False),
    *[("NDVI / dryness / moisture", imp6.loc[f_, "lgbm_all"], 3, False) for f_ in ("NDVI", "NDBI", "MNDWI")],
    *[("grass / tree / built-up", imp6.loc[f_, "lgbm_all"], 3, False) for f_ in ("frac_grass", "frac_tree", "frac_built")],
    ("Spectral block (summed", gs6.loc[("lgbm", "all"), "group_mean_abs"], 3, False),
    ("cancel within the block", 100 * gs6.loc[("lgbm", "all"), "cancellation_share"], 0, False),
    ("rank agreement (Spearman)", ra.loc["lgbm vs xgb, all cells", "spearman"], 3, False),
    ("LightGBM vs XGBoost: group", 100 * ds.loc["group", "lgbm_vs_xgb"], 1, False),
    ("LightGBM vs XGBoost: group", 100 * ds.loc["feature", "lgbm_vs_xgb"], 1, False),
    ("interventional vs path-dependent", 100 * ds.loc["group", "interventional_vs_path_dependent"], 1, False),
    ("interventional vs path-dependent", 100 * ds.loc["feature", "interventional_vs_path_dependent"], 1, False),
    ("NDVI-dryness correlation", ic8[(ic8.a == "NDVI") & (ic8.b == "NDBI")].r.iloc[0], 2, False),
    ("Tree-cover knot", 100 * t.pos_knot, 1, False), ("Tree-cover knot", 100 * t.pos_knot_lo, 1, False),
    ("Tree-cover knot", 100 * t.pos_knot_hi, 1, False),
    ("step at first canopy", t.step_zero_to_low, 2, False),
    ("below / above knot", t.pos_slope_below / 10, 2, False), ("below / above knot", t.pos_slope_above / 10, 2, False),
    ("Moisture-index onset", thr.loc["MNDWI", "knot"], 2, False),
    ("raw vs model", rq5.loc["all", "raw_slope_per_0.1"], 2, False),
    ("raw vs model", rq5.loc["all", "shap_slope_per_0.1"], 3, False),
    ("Shrinkage of the built-up", 100 * (1 - rq5.loc["all", "shap_slope_per_0.1"] / rq5.loc["all", "raw_slope_per_0.1"]), 0, False),
    ("within urban cells", rq5.loc["urban", "raw_pearson"], 3, False),
    ("Target cells (hottest", len(tc), 0, True),
    ("First point of canopy", ms.loc["realistic", "first_point_0_to_1pct"], 2, False),
    ("First point of canopy", ms.loc["lower", "first_point_0_to_1pct"], 2, False),
    ("Per point, realistic", ms.loc["realistic", "per_point_2pct_to_knot"], 2, False),
    ("Per point, realistic", ms.loc["realistic", "per_point_12_to_30pct"], 2, False),
    *[("Index-model R²", cv7[i_], 3, False) for i_ in ("MNDWI", "NDVI", "NDBI")],
    ("outside training support", ens.pct_flagged.max(), 2, False),
]
for kw_, val_, nd_, th_ in VIVA:
    check(kw_, kw_, val_, nd_, th_, doc=V)
for row_, sc_ in (("A +10 points", "A"), ("B +20 points", "B"), ("C to the knot", "C")):
    for b_ in ("lower", "realistic"):
        check(f"{sc_} {b_}", row_, ens.loc[(sc_, b_), "dLST_excl_first_point"], 2, doc=V)
check_findings("targets with no tree cover", "Targets with no tree cover", "78.7%", "phase7", doc=V)
check("Moran's I about", "Moran's I of about", lg.mean(), 1, doc=V)

# ===========================================================================
# Coverage guard: every number in the deck and the viva notes is checked
# ===========================================================================
from mirageai import config as cfg  # noqa: E402
from mirageai import explain as ex  # noqa: E402
from mirageai import models as mo  # noqa: E402
from mirageai import scenarios as sc  # noqa: E402

CONSTANTS = {
    str(cfg.SCALE),                                              # 100 m grid
    *(str(y) for y in cfg.YEARS), "2021",                        # study years; WorldCover epoch (D-006)
    *(str(b // 1000) for b in cfg.BLOCK_SIZES_M),                # 1, 2, 4, 8 km blocks
    str(len(cfg.CV_SEEDS)), str(cfg.N_SPLITS),                   # 5 seeds, 5 folds
    str(len(cfg.FEATURES_TREE)), str(len(cfg.FRACTION_FEATURES)), str(len(cfg.YEARS)),
    *(f"{100 * float(sc.scenario_delta(s_, [0.0], 0.0)[0]):.0f}" for s_ in ("A", "B")),  # +10 / +20 points
    f"{100 * t.pos_knot:.0f}",                                   # "~10%" canopy (knot, D-037)
    f"{100 * t.pos_knot + 2:.0f}",                               # "~10-12%": the 12-30% band edge (D-040)
    f"{ex.N_EXPLAIN:,}", f"{ex.N_BACKGROUND:,}", f"{ex.N_INTERACTION:,}",
    str(mo.TUNE_N_ITER), f"{mo.TUNE_SUBSAMPLE:,}", str(sc.SUPPORT_K),
    f"{100 * sc.SUPPORT_QUANTILE:.0f}", f"{100 * cfg.URBAN_BUILT_THRESHOLD:.0f}",
    f"{100 * cfg.WATER_EXCLUDE_THRESHOLD:.0f}", str(cfg.SEED),
    *(str(s_) for s_ in cfg.CV_SEEDS),
    str(len(cfg.CV_SEEDS) * cfg.N_SPLITS),                       # 25 seed x fold runs
    str(__import__("inspect").signature(ex.breakpoint).parameters["n_boot"].default),  # bootstrap resamples
    str(cfg.NATIVE_SCALE["worldcover"]),                         # 10 m WorldCover pixel
    str(cfg.NATIVE_SCALE["landsat"]),                            # 30 m Landsat pixel; 12-30% band edge (D-040)
    "95",                                                        # 95% bootstrap interval (explain.breakpoint, D-037)
}
STRIP = [r"\[(?:slide|notes) \d+\]", r"(?<=\| )\d(?= \|)",  # a shape that is only an ordinal label
         r"^\*\*\d+\.",                                          # viva question numbers
         r"\bD-\d+\b", r"\bT\d\.\d\b", r"\bRQ\d\b",
         r"Phase \d", r"\b20\d\d-\d\d\b", r"Landsat 8(?:/| and )9", r"Sentinel-2", r"TE7483",
         r"EPSG:\d+", r"\bR²", r"Level-2|Collection 2",
         r"\b\d{1,2}:\d\d AM\b"]                                 # overpass clock time (D-047)
for doc in ("slides", "viva"):
    for para in DOCS[doc]:
        text = para
        for pat_ in STRIP:
            text = re.sub(pat_, " ", text)
        for tok in re.findall(r"(?<![\w.])\d[\d,]*(?:\.\d+)?", text):
            tok = tok.rstrip(",")
            if tok not in COVERED[doc] and tok not in CONSTANTS:
                where_ = re.match(r"\[(slide|notes) \d+\]", para)
                results.append((False, f"[{doc}] unchecked number", tok,
                                where_.group(0) if where_ else para[:40], 0))

# ---------------------------------------------------------------------------
bad = [r for r in results if not r[0]]
for ok, label, s, kw, n in results:
    if not ok:
        print(f"FLAG  {label:<32} expected '{s}' in a paragraph containing '{kw}' ({n} candidate paragraphs)")
print(f"{len(results) - len(bad)} of {len(results)} checks passed, {len(bad)} flagged")
sys.exit(1 if bad else 0)
